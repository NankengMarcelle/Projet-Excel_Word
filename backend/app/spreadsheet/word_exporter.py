import re
from datetime import date, datetime
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Emu, Inches, Pt, RGBColor
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet as OpenpyxlWorksheet

from app.spreadsheet.cell_signal import color_to_hex, used_range


def _docx_rgb(color) -> str | None:
    """color_to_hex() returns openpyxl's raw 8-char ARGB string (e.g. "00FFFF00") — the right
    shape for the JSON API (worksheet_service.py), which is what it's shared with. Word wants
    plain 6-char RGB in both places this is used: RGBColor.from_string() raises outright on
    anything but 6 hex characters (confirmed live — this exact gap crashed every conversion
    that touched a colored font with a 500, not a hang: `ValueError: RGBColor() takes three
    integer values 0-255`), and <w:shd w:fill="..."> is documented as ST_HexColor (6-digit
    RGB or "auto") — an 8-char value there isn't valid OOXML even where it doesn't outright
    crash. Strips the leading alpha byte the same way this file's own color handling always
    did, pre-refactor."""
    hex_str = color_to_hex(color)
    if hex_str and len(hex_str) == 8:
        return hex_str[2:]
    return hex_str

_HORIZONTAL_ALIGNMENT_MAP = {
    "left": WD_ALIGN_PARAGRAPH.LEFT,
    "center": WD_ALIGN_PARAGRAPH.CENTER,
    "right": WD_ALIGN_PARAGRAPH.RIGHT,
    "centerContinuous": WD_ALIGN_PARAGRAPH.CENTER,
    "justify": WD_ALIGN_PARAGRAPH.JUSTIFY,
}

_VERTICAL_ALIGNMENT_MAP = {
    "top": WD_ALIGN_VERTICAL.TOP,
    "center": WD_ALIGN_VERTICAL.CENTER,
    "bottom": WD_ALIGN_VERTICAL.BOTTOM,
}

# python-docx has no public API for cell shading/borders; both require
# reaching into the raw OOXML (<w:tcPr>) directly.


def _set_cell_shading(cell, hex_color: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_color)
    tc_pr.append(shd)


# Excel border style name -> (Word's w:val, w:sz in eighths-of-a-point). Previously every
# border rendered identically as a thin black single line regardless of the source's actual
# style or color — a report with a deliberate thick outer border and thin inner grid (a very
# common real-world layout) came out looking uniformly flat. Word's ST_Border vocabulary
# doesn't have a 1:1 match for every Excel style (there's no dash-dot-with-small-gap
# equivalent, for instance) — these are the closest available, not exact.
_BORDER_STYLE_TO_WORD: dict[str, tuple[str, int]] = {
    "hair": ("single", 2),
    "thin": ("single", 4),
    "medium": ("single", 12),
    "thick": ("single", 24),
    "double": ("double", 4),
    "dotted": ("dotted", 4),
    "dashed": ("dashed", 4),
    "dashDot": ("dotDash", 4),
    "dashDotDot": ("dotDotDash", 4),
    "mediumDashed": ("dashed", 12),
    "mediumDashDot": ("dotDash", 12),
    "mediumDashDotDot": ("dotDotDash", 12),
    "slantDashDot": ("dashed", 12),
}
_DEFAULT_BORDER_COLOR = "000000"


def _set_cell_borders(cell, sides: dict) -> None:
    """`sides` maps edge name -> openpyxl `Side` object (or None/styleless), not just a
    presence flag — so the actual style and color make it through instead of every border
    being forced to the same thin black line."""
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_borders = OxmlElement("w:tcBorders")
    for edge_name, side in sides.items():
        if not side or not side.style:
            continue
        word_val, size = _BORDER_STYLE_TO_WORD.get(side.style, ("single", 4))
        color = _docx_rgb(getattr(side, "color", None)) or _DEFAULT_BORDER_COLOR
        edge_el = OxmlElement(f"w:{edge_name}")
        edge_el.set(qn("w:val"), word_val)
        edge_el.set(qn("w:sz"), str(size))
        edge_el.set(qn("w:color"), color)
        tc_borders.append(edge_el)
    tc_pr.append(tc_borders)


# Excel's own default column width (in its character-count unit) when a column has no
# explicit <col> dimension at all — openpyxl leaves ws.column_dimensions empty for such
# columns rather than reporting this default, so callers have to supply it themselves.
_DEFAULT_EXCEL_COLUMN_WIDTH = 8.43

# Word's own effective default when a cell carries no explicit Excel font-size override —
# close enough to Excel's own common default (Calibri 11) that such cells shrink in step with
# the ones that do have an explicit size, instead of standing out at a fixed size.
_DEFAULT_FONT_SIZE_PT = 11.0

# However small the fit-to-page math below would like to go, still-legible text beats an
# unreadable sliver — a floor, not a target. Real many-column report sheets routinely land
# well above this (a readable ~6-8pt) once scaled; this only guards the pathological extreme.
_MIN_FONT_SIZE_PT = 6.0


def _compute_fit_to_page(ws, max_col: int, available_width: int) -> tuple[list[int], float]:
    """Mirrors Excel's own "Fit to page width" print scaling: rather than keeping every column
    at its normal Excel width and letting Word wrap (or crush) whatever doesn't fit, shrink the
    *whole layout* — column widths and font size together, by one uniform factor — so the table
    fits the page without ever wrapping text mid-word. A many-column report (this app's real
    workbooks routinely have 18-19 columns) simply can't keep a normal ~11pt font at full
    column width on one landscape page. The approach tried first — keep font size fixed, only
    shrink columns — is what caused headers like "Autorisations d'Engagement" to wrap
    letter-by-letter, confirmed live against a real workbook; rotating just the long headers
    was tried next, but that's not actually what was wanted — shrinking the font, the same way
    Excel's own print scaling does, is. Only ever shrinks text (never enlarges it): a small,
    already-fitting sheet's columns still expand to fill the page, but its font stays as-is."""
    raw_widths = []
    for col_index in range(1, max_col + 1):
        dimension = ws.column_dimensions.get(get_column_letter(col_index))
        excel_width = dimension.width if dimension and dimension.width else _DEFAULT_EXCEL_COLUMN_WIDTH
        raw_widths.append(Inches(excel_width / 7.0))
    total = sum(raw_widths, Emu(0))
    width_scale = available_width / total
    column_widths = [Emu(int(width * width_scale)) for width in raw_widths]
    font_scale = min(1.0, width_scale)
    return column_widths, font_scale


_NUMBER_FORMAT_BRACKET_RE = re.compile(r"\[[^\]]*\]")
# Excel number-format "invisible spacer"/fill escapes: `_x` reserves blank space the width of
# character x *without displaying x*, `\x` renders x as a plain literal (not part of the
# numeric pattern), and `*x` repeat-fills the column with x. All three are routinely paired
# with a currency symbol purely so plain-integer columns visually align with real currency
# columns next to them — confirmed live against a real report: `_-* #,##0\ _€_-` is an
# accounting-style *whole-number* format (no currency at all), but its literal "€" character
# was being read as "this is currency" and prefixed onto every value, which Excel itself never
# actually displays. Strip these before checking for a currency symbol.
_NUMBER_FORMAT_SPACER_RE = re.compile(r"[_\\*].")
_CURRENCY_SYMBOLS = ("$", "€", "£", "¥")


def _format_number(value, number_format: str) -> str | None:
    """Best-effort rendering of Excel's number_format codes for the patterns this app's real
    (financial/report) workbooks actually use: percentages, thousands separators, fixed
    decimals, simple currency. Returns None — caller falls back to the plain value — for
    anything outside that: this is deliberately not a full Excel format-code interpreter
    (hundreds of edge cases: date/time sub-codes, per-sign sections, conditional colors,
    custom literal text), just the common numeric subset that was previously entirely
    unhandled (every percentage/currency/thousands-formatted cell showed its raw float, e.g.
    "0.15" instead of "15%")."""
    if not isinstance(value, (int, float)) or isinstance(value, bool):
        return None
    if not number_format or number_format == "General":
        return None

    # Multi-section formats ("positive;negative;zero;text") — only the first (positive)
    # section is interpreted. Strip conditional/color/locale tags like [Red] or
    # [$€-x-euro2] first; they're not part of the actual numeric pattern. Then strip
    # underscore/backslash/asterisk spacer-escapes (see _NUMBER_FORMAT_SPACER_RE) — a
    # currency symbol appearing only inside one of those is cosmetic alignment, not a real
    # currency marker.
    fmt = _NUMBER_FORMAT_BRACKET_RE.sub("", number_format).split(";")[0]
    fmt = _NUMBER_FORMAT_SPACER_RE.sub("", fmt)
    if not any(ch in fmt for ch in "0#"):
        return None  # not a recognizable numeric pattern (e.g. a custom text-only format)

    is_percent = fmt.rstrip().endswith("%")
    if is_percent:
        fmt = fmt.rstrip()[:-1]
        value = value * 100

    currency_symbol = next((s for s in _CURRENCY_SYMBOLS if s in fmt), None)
    if currency_symbol:
        fmt = fmt.replace(currency_symbol, "")

    use_thousands = "," in fmt
    decimal_places = len(re.findall(r"[0#]", fmt.split(".", 1)[1])) if "." in fmt else 0

    formatted = f"{value:,.{decimal_places}f}" if use_thousands else f"{value:.{decimal_places}f}"
    if currency_symbol:
        formatted = f"{currency_symbol}{formatted}"
    if is_percent:
        formatted = f"{formatted}%"
    return formatted


def _cell_display_value(cell, formula_cell=None) -> str:
    if cell.value is None:
        # `cell` comes from a data_only=True load, which for a formula cell gives *only* the
        # last cached result — nothing at all if that cache was ever lost (see
        # excel_io.save_workbook_preserving_formula_cache's own docstring for how that
        # happens). `formula_cell` is the same coordinate from a separate data_only=False
        # load, which always has the formula text regardless of cache state — falling back to
        # showing that (as Excel's own "show formulas" mode would) beats a silently blank
        # cell, even though it's the formula, not the number it would have evaluated to.
        if formula_cell is not None and formula_cell.data_type == "f":
            return str(formula_cell.value)
        return ""
    if isinstance(cell.value, (datetime, date)):
        # openpyxl hands back a real datetime/date object for a date-formatted cell (when
        # read with data_only=True) — str()'ing that directly gives an ugly
        # "2026-09-15 00:00:00" instead of anything resembling what Excel actually displays.
        # Not a full number-format-string interpreter — just the single most common case
        # that otherwise looks obviously broken.
        if isinstance(cell.value, datetime) and (cell.value.hour or cell.value.minute):
            return cell.value.strftime("%Y-%m-%d %H:%M")
        return cell.value.strftime("%Y-%m-%d")
    formatted = _format_number(cell.value, cell.number_format)
    if formatted is not None:
        return formatted
    return str(cell.value)


def worksheet_to_docx(
    ws: OpenpyxlWorksheet, output_path: Path, *, ws_formulas: OpenpyxlWorksheet | None = None
) -> None:
    """Best-effort visual mirror of a worksheet as a Word table.

    `ws` must be loaded with data_only=True (so formula cells show their calculated value,
    not their formula text). `ws_formulas` is optional: the *same* worksheet loaded with
    data_only=False instead, used only as a fallback so a formula cell whose cached value was
    lost (see excel_io.save_workbook_preserving_formula_cache) shows its formula rather than
    going silently blank — see _cell_display_value(). Without it, such cells just render
    empty, which is why every existing caller that doesn't have a second load handy (all of
    this file's own tests) still works fine passing only `ws`.

    Column widths and font size are scaled together to fit one page width (see
    _compute_fit_to_page()) — the same tradeoff Excel's own "Fit to page width" print option
    makes, rather than letting Word wrap long headers mid-word in narrow columns.

    Not pixel-perfect: Word interprets row height as a minimum (not exact), Excel's number
    format codes are only reproduced for dates and the common numeric subset (percentages,
    thousands separators, fixed decimals, simple currency — see _format_number()). Charts,
    images, and conditional formatting are not reproduced. Theme-based colors (as opposed to a
    cell's own explicit RGB) aren't resolved — see color_to_hex()'s own limitation.
    """
    document = Document()
    document.add_heading(ws.title, level=1)

    section = document.sections[0]
    # Landscape + slim margins: a wide, many-column sheet (routine for this app's real
    # workbooks) fits dramatically more per page than portrait's default ~1" margins allow —
    # the same tradeoff a manual Excel-to-Word export would reach for by hand.
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width, section.page_height = section.page_height, section.page_width
    section.top_margin = Cm(1.5)
    section.bottom_margin = Cm(1.5)
    section.left_margin = Cm(1.5)
    section.right_margin = Cm(1.5)

    # The load-bearing fix here: ws.max_row/ws.max_column reflect the sheet's *declared*
    # dimensions, which for a real, long-lived workbook can be dramatically larger than its
    # actual content (formatting once applied across a huge range and never fully cleared —
    # see used_range()'s own docstring for a concrete measured example: 277,200 declared
    # cells, 7,395 with any real content). Building and styling a table sized to the
    # declared dimensions instead of the real content is what previously made a single
    # bloated sheet take minutes to convert into a 50+ page, mostly-blank document.
    # Prefer the formulas view for this, when available: a formula cell whose cached value was
    # lost has .value == None in the values view with no other formatting to flag it as real
    # content, so cell_has_signal() there would miss it entirely — used_range() could exclude
    # a whole trailing row/column of formula cells, not just render them blank within it. A
    # formula cell's .value in the formulas view is always its (non-None) formula text,
    # regardless of cache state, so it's never missed there — matches how worksheet_service.py
    # already does this same check for the JSON read path.
    max_row, max_col = used_range(ws_formulas if ws_formulas is not None else ws)
    table = document.add_table(rows=max_row, cols=max_col)
    table.style = "Table Grid"

    # Fixed layout + explicit widths summing to exactly the page's available width, plus a
    # matching font-size shrink — see _compute_fit_to_page()'s docstring for why leaving column
    # sizing to Word's own "AutoFit to Contents" (the default) produced unusably narrow columns
    # on a wide, many-column sheet, and why shrinking columns without also shrinking the font
    # doesn't work either.
    table.autofit = False
    available_width = section.page_width - section.left_margin - section.right_margin
    column_widths, font_scale = _compute_fit_to_page(ws, max_col, available_width)
    for col_index, width in enumerate(column_widths):
        table.columns[col_index].width = width

    # The other load-bearing fix here, found live (not guessed): python-docx's `Table.cell()`
    # rebuilds the table's *entire* cell list from scratch — walking every <w:tc> element in
    # the table — on every single call (it's a plain, uncached @property under the hood, see
    # docx/table.py's `_cells`). Calling it once per cell inside this loop, as the original
    # version of this function did, made the whole export O(total_cells²): confirmed directly
    # — a used_range()-trimmed 528x19 (10,032-cell) sheet, already a 27x reduction from its
    # declared 528x525, still hadn't progressed past its first couple thousand cells after a
    # full minute. `_Row.cells` is much cheaper (it only walks that one row's own elements),
    # so fetching each row's cells exactly once up front — before any merges, which mutate
    # the tree and would invalidate cells cached across that boundary — turns this back into
    # straightforward O(total_cells) work.
    doc_rows = [row.cells for row in table.rows]

    for row_idx, row in enumerate(ws.iter_rows(min_row=1, max_row=max_row, max_col=max_col), start=1):
        doc_row_cells = doc_rows[row_idx - 1]
        for cell in row:
            doc_cell = doc_row_cells[cell.column - 1]
            formula_cell = ws_formulas.cell(row=cell.row, column=cell.column) if ws_formulas is not None else None
            doc_cell.text = _cell_display_value(cell, formula_cell)
            paragraph = doc_cell.paragraphs[0]
            run = paragraph.runs[0] if paragraph.runs else paragraph.add_run("")

            font = cell.font
            run.bold = bool(font.bold)
            run.italic = bool(font.italic)
            run.underline = bool(font.underline and font.underline != "none")
            if font.name:
                run.font.name = font.name
            base_size = font.size or _DEFAULT_FONT_SIZE_PT
            run.font.size = Pt(max(_MIN_FONT_SIZE_PT, base_size * font_scale))
            font_color = _docx_rgb(font.color)
            if font_color:
                run.font.color.rgb = RGBColor.from_string(font_color)

            alignment = cell.alignment
            if alignment.horizontal in _HORIZONTAL_ALIGNMENT_MAP:
                paragraph.alignment = _HORIZONTAL_ALIGNMENT_MAP[alignment.horizontal]
            if alignment.vertical in _VERTICAL_ALIGNMENT_MAP:
                doc_cell.vertical_alignment = _VERTICAL_ALIGNMENT_MAP[alignment.vertical]

            # Only "solid" actually paints fgColor as a flat background in Excel's own
            # rendering — any other fill_type (most commonly "gray125", the legacy marker
            # OOXML silently writes onto a cell merely touched by formatting without an
            # explicit fill) would otherwise get its often grey/black fgColor painted as a
            # real background here, which is not what the workbook shows on screen. Same
            # rule the JSON read path already applies (worksheet_service.py).
            if cell.fill and cell.fill.fill_type == "solid":
                fill_color = _docx_rgb(cell.fill.fgColor)
                if fill_color:
                    _set_cell_shading(doc_cell, fill_color)

            border = cell.border
            sides = {
                "top": border.top,
                "bottom": border.bottom,
                "left": border.left,
                "right": border.right,
            }
            if any(side and side.style for side in sides.values()):
                _set_cell_borders(doc_cell, sides)

    for merged_range in ws.merged_cells.ranges:
        # A merge can extend past the trimmed used_range() boundary only if its own anchor
        # cell had no real signal of its own (rare — merging is itself a form of formatting
        # cell_has_signal would normally catch via the anchor's border/fill/alignment) — skip
        # rather than raise on a stale/out-of-range merge definition.
        if merged_range.min_row > max_row or merged_range.min_col > max_col:
            continue
        end_row = min(merged_range.max_row, max_row)
        end_col = min(merged_range.max_col, max_col)
        # table.rows[r].cells[c] here, not table.cell(r, c): each merge can change the tree
        # (so, unlike the main loop above, these lookups can't be cached once up front — they
        # need to be live/current every time), but table.cell() pays for that by rebuilding
        # the *entire* table's cell list on every call, same O(total_cells) cost as before,
        # now paid twice per merge. table.rows[r].cells only walks that one row's own
        # elements. Confirmed live on the real sheet this was found on: 257 merges through
        # table.cell() was the dominant cost of a 343-second conversion — this row-scoped
        # version is still always correct (never stale) but doesn't pay to re-walk every
        # other row's cells just to find the two this particular merge needs.
        start_cell = table.rows[merged_range.min_row - 1].cells[merged_range.min_col - 1]
        end_cell = table.rows[end_row - 1].cells[end_col - 1]
        if start_cell is end_cell:
            continue
        # merge() concatenates *every* absorbed cell's own paragraphs onto the anchor's —
        # and every cell in this table starts with one (usually empty) paragraph, even ones
        # never written to. So merging N cells leaves N-1 blank trailing paragraphs stacked
        # after the anchor's real text. Confirmed live, not guessed: a single-row, 17-column
        # title merge — one tight line in the source — came out as real text followed by 7+
        # blank lines in Word, a large visible gap that looked like a row-height or wrapping
        # bug but wasn't. Record the anchor's own (legitimate) paragraph count *before*
        # merging, then trim the merged result back down to exactly that — not just "keep the
        # first paragraph": a cell whose own source value has a real embedded newline (this
        # sheet's title row does) legitimately has more than one paragraph of its own, and
        # that has to survive the trim.
        original_paragraph_count = len(start_cell.paragraphs)
        merged_cell = start_cell.merge(end_cell)
        for extra_paragraph in merged_cell.paragraphs[original_paragraph_count:]:
            extra_paragraph._element.getparent().remove(extra_paragraph._element)

    # Word treats row height as a minimum, not exact — this still helps rows with
    # deliberately tall content (wrapped text, larger fonts) come out closer to Excel's
    # proportions instead of Word's own default single-line row height. Scaled by the same
    # font_scale as the text itself, so a shrunk font doesn't leave rows far taller than the
    # (now smaller) text actually needs.
    for row_index in range(1, max_row + 1):
        dimension = ws.row_dimensions.get(row_index)
        if dimension and dimension.height:
            table.rows[row_index - 1].height = Pt(dimension.height * font_scale)

    output_path.parent.mkdir(parents=True, exist_ok=True)
    document.save(output_path)
