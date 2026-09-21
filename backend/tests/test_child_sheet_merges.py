import io

from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook

from app.spreadsheet.filter_engine import compute_projected_merges


# --- Unit tests: compute_projected_merges() in isolation --------------------------------
#
# The function only reads ws.merged_cells.ranges, so a bare Workbook with nothing but
# ws.merge_cells() calls is enough — no need to route through the API/DB for these.


def test_header_merge_survives_column_reorder():
    ws = Workbook().active
    ws.merge_cells("B1:D1")  # a group header spanning original columns 2-4

    merges = compute_projected_merges(
        ws, header_start_row=1, header_end_row=1, selected_columns=[4, 2, 3, 1], row_mask=[]
    )

    # col 2 -> position 2, col 3 -> position 3, col 4 -> position 1: surviving span is 2-4... no,
    # min/max of {1, 2, 3} = 1..3.
    assert merges == [(1, 1, 1, 3)]


def test_header_merge_shrinks_when_one_of_its_columns_is_dropped():
    ws = Workbook().active
    ws.merge_cells("B1:D1")  # original columns 2-4

    merges = compute_projected_merges(
        ws, header_start_row=1, header_end_row=1, selected_columns=[1, 2, 4], row_mask=[]
    )

    # Column 3 dropped; columns 2 and 4 survive at positions 2 and 3.
    assert merges == [(1, 1, 2, 3)]


def test_header_merge_dropped_when_it_collapses_to_a_single_column():
    ws = Workbook().active
    ws.merge_cells("B1:D1")  # original columns 2-4

    merges = compute_projected_merges(
        ws, header_start_row=1, header_end_row=1, selected_columns=[1, 2], row_mask=[]
    )

    assert merges == []


def test_data_merge_remaps_rows_after_filtering():
    ws = Workbook().active
    ws.merge_cells("A3:A5")  # a vertical data-label merge, original rows 3-5

    # header_end_row=1 -> data rows start at 2: row2=idx0, row3=idx1, row4=idx2, row5=idx3.
    merges = compute_projected_merges(
        ws,
        header_start_row=1,
        header_end_row=1,
        selected_columns=[1],
        row_mask=[True, True, True, True],
    )

    # idx1/2/3 survive at output positions 2/3/4; header block is 1 row, so output rows 3-5.
    assert merges == [(3, 5, 1, 1)]


def test_data_merge_with_non_contiguous_survivors_spans_their_new_positions():
    ws = Workbook().active
    ws.merge_cells("A3:A5")  # original rows 3-5 -> idx1, idx2, idx3

    # idx2 (original row 4) is filtered out; idx1 and idx3 survive.
    merges = compute_projected_merges(
        ws,
        header_start_row=1,
        header_end_row=1,
        selected_columns=[1],
        row_mask=[True, True, False, True],
    )

    # idx0 -> position 1, idx1 -> position 2, idx3 -> position 3 (idx2 has no position at all).
    # Surviving positions for this merge: {2, 3} -> output rows 3-4.
    assert merges == [(3, 4, 1, 1)]


def test_data_merge_dropped_when_every_one_of_its_rows_is_filtered_out():
    ws = Workbook().active
    ws.merge_cells("A3:A4")  # original rows 3-4 -> idx1, idx2

    merges = compute_projected_merges(
        ws,
        header_start_row=1,
        header_end_row=1,
        selected_columns=[1],
        row_mask=[True, False, False, True],
    )

    assert merges == []


def test_merge_spanning_the_header_data_boundary_is_dropped():
    ws = Workbook().active
    ws.merge_cells("A1:A2")  # row 1 is header, row 2 is already data (header_end_row=1)

    merges = compute_projected_merges(
        ws, header_start_row=1, header_end_row=1, selected_columns=[1], row_mask=[True]
    )

    assert merges == []


# --- End-to-end: create/sync a child sheet through the real API and inspect the file ------
#
# Regression coverage for the reported bug: "the merged data columns were unmerged and now
# everything is choked up" — read_rows()'s merge-anchor resolution already propagates a
# merge's *value* to every cell it covers, but write_rows() never recreated the merge
# *structure* itself, so every generated/synced child sheet came out fully unmerged.


def _build_merged_workbook_bytes() -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Matrix"

    # Two-row header: "Action" (col 1) is vertically merged across both header rows (no
    # sub-header of its own); "Totals" (cols 2-3) is horizontally merged across row 1, with
    # its own leaf labels ("AE", "CP") on row 2.
    ws.cell(row=1, column=1, value="Action")
    ws.merge_cells("A1:A2")
    ws.cell(row=1, column=2, value="Totals")
    ws.merge_cells("B1:C1")
    ws.cell(row=2, column=2, value="AE")
    ws.cell(row=2, column=3, value="CP")

    # A vertically-merged data label ("Group X") spanning two task rows, and an unmerged
    # single row ("Group Y") right after.
    ws.cell(row=3, column=1, value="Group X")
    ws.merge_cells("A3:A4")
    ws.cell(row=3, column=2, value=10)
    ws.cell(row=3, column=3, value=20)
    ws.cell(row=4, column=2, value=30)
    ws.cell(row=4, column=3, value=40)
    ws.cell(row=5, column=1, value="Group Y")
    ws.cell(row=5, column=2, value=50)
    ws.cell(row=5, column=3, value=60)

    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def _upload(api_client: TestClient, headers: dict, xlsx_bytes: bytes) -> tuple[str, str]:
    response = api_client.post(
        "/workbooks",
        headers=headers,
        files={"file": ("matrix.xlsx", xlsx_bytes, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
    )
    assert response.status_code == 201
    workbook_id = response.json()["id"]
    parent_worksheet_id = api_client.get(f"/workbooks/{workbook_id}", headers=headers).json()["worksheets"][0]["id"]
    return workbook_id, parent_worksheet_id


def _download_workbook(api_client: TestClient, headers: dict, workbook_id: str):
    response = api_client.get(f"/workbooks/{workbook_id}/download", headers=headers)
    assert response.status_code == 200
    return load_workbook(io.BytesIO(response.content))


def _create_child_sheet(api_client: TestClient, headers: dict, workbook_id: str, *, child_sheet_name: str, sources: list[dict]):
    return api_client.post(
        f"/workbooks/{workbook_id}/child-sheets",
        headers=headers,
        json={"child_sheet_name": child_sheet_name, "sources": sources},
    )


def test_create_child_sheet_preserves_merges_with_no_column_reorder(
    api_client: TestClient, auth_headers: dict
):
    workbook_id, parent_worksheet_id = _upload(api_client, auth_headers, _build_merged_workbook_bytes())

    response = _create_child_sheet(
        api_client,
        auth_headers,
        workbook_id,
        child_sheet_name="Matrix Copy",
        sources=[
            {
                "parent_worksheet_id": parent_worksheet_id,
                "header_start_row": 1,
                "header_end_row": 2,
                "selected_columns": [1, 2, 3],
                "filter_criteria": {"logic": "AND", "conditions": []},
            }
        ],
    )
    assert response.status_code == 201

    wb = _download_workbook(api_client, auth_headers, workbook_id)
    child_ws = wb["Matrix Copy"]
    merge_strings = {str(r) for r in child_ws.merged_cells.ranges}

    assert "A1:A2" in merge_strings  # "Action" vertical header merge, unchanged
    assert "B1:C1" in merge_strings  # "Totals" horizontal header merge, unchanged
    assert "A3:A4" in merge_strings  # "Group X" vertical data-label merge, unchanged
    assert child_ws["A1"].value == "Action"
    assert child_ws["B1"].value == "Totals"
    assert child_ws["A3"].value == "Group X"


def test_create_child_sheet_remaps_merges_for_column_reorder_and_drop(
    api_client: TestClient, auth_headers: dict
):
    workbook_id, parent_worksheet_id = _upload(api_client, auth_headers, _build_merged_workbook_bytes())

    # Drop column 2 ("AE") and put column 3 ("CP") before column 1 ("Action").
    response = _create_child_sheet(
        api_client,
        auth_headers,
        workbook_id,
        child_sheet_name="Matrix Reordered",
        sources=[
            {
                "parent_worksheet_id": parent_worksheet_id,
                "header_start_row": 1,
                "header_end_row": 2,
                "selected_columns": [3, 1],
                "filter_criteria": {"logic": "AND", "conditions": []},
            }
        ],
    )
    assert response.status_code == 201

    wb = _download_workbook(api_client, auth_headers, workbook_id)
    child_ws = wb["Matrix Reordered"]
    merge_strings = {str(r) for r in child_ws.merged_cells.ranges}

    # "Totals" (originally B1:C1) collapses to a single surviving column (CP, now column 1)
    # and is correctly dropped rather than kept as a degenerate 1x1 "merge" — no merge should
    # span row 1 alone at all.
    assert not any(cell_range.min_row == cell_range.max_row == 1 for cell_range in child_ws.merged_cells.ranges)

    # "Action" (originally col 1, now col 2) is a single-column vertical merge either way —
    # column reordering doesn't touch its row span, just its column position.
    assert "B1:B2" in merge_strings
    assert child_ws["B1"].value == "Action"

    # "Group X" (originally col 1 rows 3-4, now col 2) survives the same way.
    assert "B3:B4" in merge_strings
    assert child_ws["B3"].value == "Group X"


def test_sync_child_sheet_recreates_merges_and_clears_stale_ones(
    api_client: TestClient, auth_headers: dict
):
    # Regression test for write_rows()'s own new behavior: on a re-sync, any merges from the
    # *previous* generation of this same child sheet must be cleared before the new ones are
    # applied, or a stale range left over from before would collide with (or misdescribe)
    # freshly written data.
    workbook_id, parent_worksheet_id = _upload(api_client, auth_headers, _build_merged_workbook_bytes())

    create_response = _create_child_sheet(
        api_client,
        auth_headers,
        workbook_id,
        child_sheet_name="Matrix Copy",
        sources=[
            {
                "parent_worksheet_id": parent_worksheet_id,
                "header_start_row": 1,
                "header_end_row": 2,
                "selected_columns": [1, 2, 3],
                "filter_criteria": {"logic": "AND", "conditions": []},
            }
        ],
    )
    child_worksheet_id = create_response.json()["worksheet"]["id"]

    # Force a real re-sync by editing the parent first.
    api_client.put(
        f"/workbooks/{workbook_id}/worksheets/{parent_worksheet_id}",
        headers=auth_headers,
        json={"edits": [{"row": 5, "column": 2, "value": 99}]},
    )
    sync_response = api_client.post(
        f"/workbooks/{workbook_id}/child-sheets/by-child/{child_worksheet_id}/sync", headers=auth_headers
    )
    assert sync_response.status_code == 200

    wb = _download_workbook(api_client, auth_headers, workbook_id)
    child_ws = wb["Matrix Copy"]
    merge_strings = [str(r) for r in child_ws.merged_cells.ranges]

    # Exactly one of each merge — not duplicated, not stale-plus-fresh.
    assert merge_strings.count("A1:A2") == 1
    assert merge_strings.count("B1:C1") == 1
    assert merge_strings.count("A3:A4") == 1
    assert child_ws["B5"].value == 99
