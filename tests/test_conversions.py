import io

from docx import Document
from fastapi.testclient import TestClient


def _upload_sample(api_client: TestClient, headers: dict, sample_xlsx_bytes: bytes) -> dict:
    response = api_client.post(
        "/workbooks",
        headers=headers,
        files={"file": ("sample.xlsx", sample_xlsx_bytes, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
    )
    assert response.status_code == 201
    return response.json()


def test_convert_worksheet_and_download_once(
    api_client: TestClient, auth_headers: dict, sample_xlsx_bytes: bytes
):
    created = _upload_sample(api_client, auth_headers, sample_xlsx_bytes)
    workbook_id = created["id"]
    worksheet_id = api_client.get(
        f"/workbooks/{workbook_id}", headers=auth_headers
    ).json()["worksheets"][0]["id"]

    convert_response = api_client.post(f"/worksheets/{worksheet_id}/convert", headers=auth_headers)
    assert convert_response.status_code == 201
    body = convert_response.json()
    conversion_id = body["conversion"]["id"]
    assert body["word_document"]["filename"] == "Data.docx"

    meta_response = api_client.get(f"/conversions/{conversion_id}", headers=auth_headers)
    assert meta_response.status_code == 200

    download_response = api_client.get(f"/conversions/{conversion_id}/download", headers=auth_headers)
    assert download_response.status_code == 200

    document = Document(io.BytesIO(download_response.content))
    table = document.tables[0]
    header_texts = [cell.text for cell in table.rows[0].cells]
    assert header_texts == ["Name", "Status", "Amount"]
    assert table.rows[0].cells[0].paragraphs[0].runs[0].bold is True

    second_download = api_client.get(f"/conversions/{conversion_id}/download", headers=auth_headers)
    assert second_download.status_code == 410


def test_convert_worksheet_requires_ownership(
    api_client: TestClient, auth_headers: dict, sample_xlsx_bytes: bytes
):
    from tests.conftest import register_and_login

    created = _upload_sample(api_client, auth_headers, sample_xlsx_bytes)
    workbook_id = created["id"]
    worksheet_id = api_client.get(
        f"/workbooks/{workbook_id}", headers=auth_headers
    ).json()["worksheets"][0]["id"]

    _, other_token = register_and_login(api_client)
    other_headers = {"Authorization": f"Bearer {other_token}"}

    response = api_client.post(f"/worksheets/{worksheet_id}/convert", headers=other_headers)
    assert response.status_code == 404
