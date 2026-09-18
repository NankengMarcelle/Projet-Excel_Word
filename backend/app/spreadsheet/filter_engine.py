from copy import copy

from openpyxl.worksheet.cell_range import CellRange
from openpyxl.worksheet.worksheet import Worksheet as OpenpyxlWorksheet

from app.spreadsheet.cell_signal import used_range


def safe_unmerge(ws: OpenpyxlWorksheet, coord: str) -> None:
    """Same as ws.unmerge_cells(coord), except tolerant of a non-anchor cell that was never
    actually present in ws._cells. openpyxl's own unmerge_cells() unconditionally does `del
    self._cells[(row, col)]` for every non-anchor cell in the merge — but a merge read from a
    real Excel-authored file can cover a cell that never had an XML <c> entry at all (a
    genuinely empty cell inside the merge), which throws a bare KeyError. Confirmed live
    against real workbook data; see worksheet_service.apply_structural_edit's own comment for
    the original discovery."""
    cell_range = CellRange(coord)
    if cell_range.coord in ws.merged_cells:
        ws.merged_cells.remove(cell_range)
    cells = cell_range.cells
    next(cells)  # skip the anchor cell, exactly like openpyxl's own unmerge_cells does
    for row, col in cells:
        ws._cells.pop((row, col), None)


def _build_merge_lookup(ws: OpenpyxlWorksheet) -> dict[tuple[int, int], tuple[int, int]]:
    """Maps every (row, col) inside a merged range to that merge's anchor (top-left) cell."""
    lookup: dict[tuple[int, int], tuple[int, int]] = {}
    for merged_range in ws.merged_cells.ranges:
        anchor = (merged_range.min_row, merged_range.min_col)
        for row in range(merged_range.min_row, merged_range.max_row + 1):
            for col in range(merged_range.min_col, merged_range.max_col + 1):
                lookup[(row, col)] = anchor
    return lookup


def _resolve_value(ws: OpenpyxlWorksheet, merge_lookup: dict, row: int, col: int):
    """A merged cell's value only lives in its top-left (anchor) cell — every other cell in
    the merge is genuinely empty in the underlying file, even though Excel visually shows the
    anchor's value across the whole merged block. Reading a cell's *own* value directly would
    make a non-anchor cell in a real matrix's merged column (a "Structure" label spanning
    several task rows, a "Prévision 2026" group header spanning several columns) look blank
    when it isn't. Resolves to the anchor's value for any merged cell, its own value
    otherwise."""
    anchor = merge_lookup.get((row, col))
    anchor_row, anchor_col = anchor if anchor is not None else (row, col)
    return ws.cell(row=anchor_row, column=anchor_col).value


def _resolve_cell(ws: OpenpyxlWorksheet, merge_lookup: dict, row: int, col: int):
    """Same merge-anchor resolution as _resolve_value, but hands back the anchor Cell itself
    rather than its value — a merge's non-anchor cells carry no formatting of their own worth
    reading, so copying a cell's *style* elsewhere needs the same anchor redirect its value
    already gets."""
    anchor = merge_lookup.get((row, col))
    anchor_row, anchor_col = anchor if anchor is not None else (row, col)
    return ws.cell(row=anchor_row, column=anchor_col)


def read_rows(
    ws: OpenpyxlWorksheet,
    header_start_row: int,
    header_end_row: int,
    *,
    bounds: tuple[int, int] | None = None,
) -> tuple[list[list], list[dict[int, object]]]:
    """Reads a worksheet with a (possibly multi-row) header block, honoring merged cells.

    Returns `(header_grid, data_rows)`:
    - `header_grid`: one list per header row (`header_start_row` through `header_end_row`
      inclusive), each holding every column's resolved value, 1-indexed (`header_grid[0][0]`
      is row `header_start_row`, column 1).
    - `data_rows`: one `{column_index: resolved_value}` dict per row after the header block.

    Column identity throughout this module is the 1-indexed column *number*, never header
    text — a real multi-row-header matrix can have the same leaf label ("AE voté") repeated
    under multiple group headers ("Prévision 2026", "Prévision 2027"), so header text can't
    safely be used as a unique key. Header text is only ever a display label for callers.

    Bounded by `used_range()`, not `ws.max_row`/`ws.max_column` — see that function's own
    docstring for why a long-lived real workbook's declared dimensions can be dramatically
    larger than any actual content.

    `bounds` lets a caller pass in a pre-computed `(max_row, max_col)` instead of deriving one
    from `ws` here — needed when this is called against a `data_only=True` view: a formula
    cell with no cached result yet reads as `None` there (vs. its always-present formula text
    in a `data_only=False` view), which can make `used_range()` under-report the sheet's real
    extent. A caller that also needs `read_row_styles()` against the *other* view of the same
    sheet must pass the same `bounds` to both, computed from whichever view is authoritative
    (the formulas view — see `used_range()`'s own docstring on the Word exporter hitting this
    exact mismatch), or the two outputs silently stop lining up row-for-row.
    """
    max_row, max_col = bounds if bounds is not None else used_range(ws)
    merge_lookup = _build_merge_lookup(ws)

    header_grid = [
        [_resolve_value(ws, merge_lookup, row, col) for col in range(1, max_col + 1)]
        for row in range(header_start_row, header_end_row + 1)
    ]
    data_rows = [
        {col: _resolve_value(ws, merge_lookup, row, col) for col in range(1, max_col + 1)}
        for row in range(header_end_row + 1, max_row + 1)
    ]
    return header_grid, data_rows


def read_row_styles(
    ws: OpenpyxlWorksheet,
    header_start_row: int,
    header_end_row: int,
    *,
    bounds: tuple[int, int] | None = None,
) -> tuple[list[list], list[dict[int, object]]]:
    """Mirrors read_rows()'s shape and merge resolution exactly, but each entry is the source
    Cell object itself rather than its resolved value — write_rows() reads
    font/fill/border/alignment/number_format off these to carry a parent sheet's formatting
    into a child sheet, which read_rows()'s plain values alone can't express. See read_rows()'s
    own docstring for why `bounds` needs to be the *same* value passed to both when reading the
    same sheet's values and styles from two different (`data_only=True`/`False`) views."""
    max_row, max_col = bounds if bounds is not None else used_range(ws)
    merge_lookup = _build_merge_lookup(ws)

    header_style_grid = [
        [_resolve_cell(ws, merge_lookup, row, col) for col in range(1, max_col + 1)]
        for row in range(header_start_row, header_end_row + 1)
    ]
    data_style_rows = [
        {col: _resolve_cell(ws, merge_lookup, row, col) for col in range(1, max_col + 1)}
        for row in range(header_end_row + 1, max_row + 1)
    ]
    return header_style_grid, data_style_rows


def apply_value_overrides(
    header_grid: list[list],
    rows: list[dict[int, object]],
    header_start_row: int,
    header_end_row: int,
    overrides: dict[tuple[int, int], object],
) -> None:
    """Patches specific cells' resolved values in place, after read_rows() but before
    filtering/projection — used to replace openpyxl's `data_only=True` cache for a formula
    cell with its *live* value from Univer's own client-side formula engine, sent by the
    frontend. openpyxl has no formula engine of its own; a formula's cached result can be
    stale or (as confirmed live against a real workbook — every "Sous-total" row's SUM
    formulas) entirely missing, even though the *live, open* spreadsheet computes and shows a
    real number just fine. `overrides` keys are 1-indexed (row, column) tuples in the parent
    sheet's own coordinate space, matching read_rows()'s own numbering — a no-op (`{}`) when
    the caller has nothing to patch.
    """
    if not overrides:
        return
    for header_row_index, header_row in enumerate(header_grid):
        actual_row = header_start_row + header_row_index
        for col_index in range(len(header_row)):
            key = (actual_row, col_index + 1)
            if key in overrides:
                header_row[col_index] = overrides[key]
    for row_index, row in enumerate(rows):
        actual_row = header_end_row + 1 + row_index
        for col in row:
            key = (actual_row, col)
            if key in overrides:
                row[col] = overrides[key]


def compute_projected_merges(
    ws: OpenpyxlWorksheet,
    header_start_row: int,
    header_end_row: int,
    selected_columns: list[int],
    row_mask: list[bool],
) -> list[tuple[int, int, int, int]]:
    """Determines which of the parent sheet's merged ranges should be recreated in a
    generated/synced child sheet, and where — remapped for both column selection/reordering
    and row filtering, the same way a real merge would look if you manually deleted the same
    rows/columns in Excel and dragged the survivors together. Ported from a colleague's
    Fortune-sheet-based generator (`calculateProjectedSheet` in
    Frontend_Net/Projet-Excel_Word/frontend/src/components/FortuneSheetEditor.jsx), whose
    child sheets preserve merges correctly, adapted to this app's own column-identity model
    (1-indexed column *number*, never header text — see read_rows()'s own docstring) and
    header/data split.

    Returns `(min_row, max_row, min_col, max_col)` tuples, all 1-indexed, already in the
    *output* sheet's coordinate space — matching exactly what write_rows() is about to write,
    so a caller just hands this straight to write_rows()'s `merges` param. A header merge's
    rows are renumbered against `header_start_row` (the header block is never filtered, only
    column-projected, so header row order/count never changes); a data merge's rows are
    renumbered against however many of `row_mask`'s True rows preceded it.

    A merge is dropped entirely if every one of its columns was excluded from
    `selected_columns`, or every one of its rows was filtered out by `row_mask`. A merge that
    survives but collapses to a single row *and* single column is also dropped — a 1x1 "merge"
    isn't a merge, and openpyxl's own `ws.merge_cells()` rejects one anyway. A merge spanning
    across the header/data boundary is dropped too (not expected in a real header block, and
    read_rows()/write_rows() already treat the two blocks as structurally separate).

    Note on non-contiguous survivors: if row filtering keeps some but not all of a vertical
    merge's original rows (e.g. rows 1 and 3 of an original 3-row merge, row 2 filtered out),
    this takes the min/max of the *surviving* rows' new positions as the merge's new span —
    same approximation the colleague's own algorithm makes for its equivalent case, not a gap
    introduced here.
    """
    col_position = {column: index for index, column in enumerate(selected_columns, start=1)}

    header_row_count = header_end_row - header_start_row + 1
    # Maps each ORIGINAL data row's 0-index (within row_mask, i.e. read_rows()'s own `rows`
    # list) to its 1-indexed position among the *surviving* rows — i.e. where it lands in the
    # output, right after the header block.
    data_row_position: dict[int, int] = {}
    next_position = 1
    for original_index, keep in enumerate(row_mask):
        if keep:
            data_row_position[original_index] = next_position
            next_position += 1

    projected: list[tuple[int, int, int, int]] = []
    for merged_range in ws.merged_cells.ranges:
        min_row, max_row = merged_range.min_row, merged_range.max_row
        min_col, max_col = merged_range.min_col, merged_range.max_col

        surviving_cols = sorted(
            col_position[col] for col in range(min_col, max_col + 1) if col in col_position
        )
        if not surviving_cols:
            continue

        if header_start_row <= min_row and max_row <= header_end_row:
            new_min_row = min_row - header_start_row + 1
            new_max_row = max_row - header_start_row + 1
        elif min_row > header_end_row:
            surviving_rows = sorted(
                data_row_position[row - header_end_row - 1]
                for row in range(min_row, max_row + 1)
                if (row - header_end_row - 1) in data_row_position
            )
            if not surviving_rows:
                continue
            new_min_row = header_row_count + surviving_rows[0]
            new_max_row = header_row_count + surviving_rows[-1]
        else:
            continue

        new_min_col, new_max_col = surviving_cols[0], surviving_cols[-1]
        if new_min_row == new_max_row and new_min_col == new_max_col:
            continue
        projected.append((new_min_row, new_max_row, new_min_col, new_max_col))

    return projected


def _evaluate_condition(condition: dict, row: dict[int, object]) -> bool:
    column = condition["column"]
    operator = condition["operator"]
    value = condition.get("value")
    cell_value = row.get(column)

    if operator == "equals":
        return cell_value == value
    if operator == "not_equals":
        return cell_value != value
    if operator == "contains":
        return cell_value is not None and value is not None and str(value) in str(cell_value)
    if operator == "greater_than":
        return cell_value is not None and cell_value > value
    if operator == "less_than":
        return cell_value is not None and cell_value < value
    if operator == "greater_or_equal":
        return cell_value is not None and cell_value >= value
    if operator == "less_or_equal":
        return cell_value is not None and cell_value <= value
    if operator == "is_empty":
        return cell_value is None or cell_value == ""
    if operator == "is_not_empty":
        return cell_value is not None and cell_value != ""
    if operator == "in":
        return cell_value in (value or [])
    raise ValueError(f"Unknown filter operator: {operator}")


def evaluate(node: dict, row: dict[int, object]) -> bool:
    """Recursively evaluate a filter_criteria node (a condition or an AND/OR group) against a
    row. Each condition's "column" is a 1-indexed column number — see read_rows()'s docstring
    for why column identity is positional, not name-based, throughout this module."""
    if "logic" in node:
        results = [evaluate(child, row) for child in node.get("conditions", [])]
        if node["logic"] == "AND":
            return all(results)
        if node["logic"] == "OR":
            return any(results)
        raise ValueError(f"Unknown filter logic: {node['logic']}")
    return _evaluate_condition(node, row)


def compute_filter_mask(rows: list[dict[int, object]], filter_criteria: dict) -> list[bool]:
    """A per-row "does this row pass the filter" boolean, in the same order as `rows` — lets a
    caller keep a second, parallel per-row list (most importantly read_row_styles()'s output,
    read from a *different* worksheet view than `rows` was) in sync with whichever rows
    apply_filter() would keep, without re-deriving row identity itself."""
    if not filter_criteria or not filter_criteria.get("conditions"):
        return [True] * len(rows)
    return [evaluate(filter_criteria, row) for row in rows]


def apply_filter(rows: list[dict[int, object]], filter_criteria: dict) -> list[dict[int, object]]:
    mask = compute_filter_mask(rows, filter_criteria)
    return [row for row, keep in zip(rows, mask) if keep]


def project_columns(rows: list[dict[int, object]], selected_columns: list[int]) -> list[list]:
    return [[row.get(column) for column in selected_columns] for row in rows]


def _copy_cell_style(target_cell, source_cell) -> None:
    """Copies the handful of formatting attributes this app's own cell-edit path already
    tracks (see CellEdit/read_worksheet_data in worksheet_service.py) — font, fill, border,
    alignment, number format — from one cell to another. `copy()` is required, not a bare
    attribute assignment: openpyxl style objects (Font, Fill, Border, Alignment) are shared,
    effectively-immutable value objects under the hood, and assigning the *same* instance to
    multiple cells across worksheets risks one later mutating in a way that silently changes
    every cell that shares it."""
    if source_cell is None:
        return
    target_cell.font = copy(source_cell.font)
    target_cell.fill = copy(source_cell.fill)
    target_cell.border = copy(source_cell.border)
    target_cell.alignment = copy(source_cell.alignment)
    target_cell.number_format = source_cell.number_format


def write_rows(
    ws: OpenpyxlWorksheet,
    header_grid: list[list],
    selected_columns: list[int],
    data_rows: list[list],
    *,
    header_style_grid: list[list] | None = None,
    data_style_rows: list[list] | None = None,
    merges: list[tuple[int, int, int, int]] | None = None,
) -> None:
    """Overwrite a worksheet's content with the header block — every header row, projected
    down to just the selected columns in the given order — followed by the projected data
    rows. Keeping the full header block (rather than flattening multiple rows into one row of
    composite names) mirrors the source matrix's own structure, the same way the original VBA
    tool this feature is modeled on did.

    `header_style_grid`/`data_style_rows` are optional, same-shape companions carrying the
    source Cell to copy formatting from for each written cell — `header_style_grid` is the
    *full*, unprojected header block (matching `header_grid`'s own shape; projected down to
    `selected_columns` here, same as the values are), `data_style_rows` is already projected
    (matching `data_rows`'s own shape). Omit both (the default) to write plain values only, as
    this always did before formatting support existed.

    `merges` (see `compute_projected_merges()`) are `(min_row, max_row, min_col, max_col)`
    1-indexed ranges, already in this call's own output coordinate space, applied via
    `ws.merge_cells()` after every value/style is written. Any merges already on `ws` (from a
    previous generation of this same child sheet, on a re-sync) are cleared first via
    `safe_unmerge()` — needed because `delete_rows()` below wipes cell content but never
    touches `ws.merged_cells` itself, so a stale range would otherwise linger and collide with
    (or just misdescribe) freshly written data. Omit `merges` (the default) to write with no
    merged cells at all, as this always did before merge-preservation existed.
    """
    for coord in [str(cell_range) for cell_range in ws.merged_cells.ranges]:
        safe_unmerge(ws, coord)

    if ws.max_row > 0:
        ws.delete_rows(1, ws.max_row)

    row_index = 1
    for header_row_index, header_row in enumerate(header_grid):
        style_row = header_style_grid[header_row_index] if header_style_grid is not None else None
        for col_index, column in enumerate(selected_columns, start=1):
            target_cell = ws.cell(row=row_index, column=col_index, value=header_row[column - 1])
            if style_row is not None:
                _copy_cell_style(target_cell, style_row[column - 1])
        row_index += 1

    for data_row_index, row in enumerate(data_rows):
        style_row = data_style_rows[data_row_index] if data_style_rows is not None else None
        for col_index, value in enumerate(row, start=1):
            target_cell = ws.cell(row=row_index, column=col_index, value=value)
            if style_row is not None:
                _copy_cell_style(target_cell, style_row[col_index - 1])
        row_index += 1

    for min_row, max_row, min_col, max_col in merges or []:
        ws.merge_cells(start_row=min_row, start_column=min_col, end_row=max_row, end_column=max_col)
