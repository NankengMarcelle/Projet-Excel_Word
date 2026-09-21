import time
import uuid
from datetime import datetime, timezone
from pathlib import Path
from typing import Callable

from fastapi import HTTPException, status
from openpyxl.utils import get_column_letter, range_boundaries
from sqlalchemy.orm import Session

from app.models.workbook import Workbook
from app.models.worksheet import Worksheet
from app.repositories import sheet_relationship_repository, worksheet_repository
from app.schemas.worksheet import CellData, WorksheetColumn, WorksheetData
from app.spreadsheet import cell_editor, excel_io, filter_engine
from app.spreadsheet.cell_signal import cell_has_signal, color_to_hex


def get_worksheet_or_404(db: Session, *, workbook_id: uuid.UUID, worksheet_id: uuid.UUID) -> Worksheet:
    worksheet = worksheet_repository.get_by_id_in_workbook(db, worksheet_id, workbook_id)
    if worksheet is None:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Worksheet not found")
    return worksheet


def read_worksheet_data(*, workbook: Workbook, worksheet: Worksheet) -> WorksheetData:
    # TEMPORARY: see apply_edits' matching timer for context — remove together.
    request_start = time.perf_counter()
    path = Path(workbook.storage_path)
    # Two separate loads: one keeps formula text, the other gives Excel's last cached
    # calculated value — openpyxl cannot return both from a single load. Both go through the
    # shared read-only cache (see excel_io.py) rather than a fresh load_workbook() per request:
    # with every worksheet in a workbook fetched in parallel when the editor opens, an
    # uncached load here meant each of those requests independently re-parsing the entire
    # file. Never close() these — the cache owns their lifecycle.
    wb_formulas = excel_io.load_workbook_cached(path, data_only=False)
    wb_values = excel_io.load_workbook_cached(path, data_only=True)
    ws_formulas = wb_formulas[worksheet.name]
    ws_values = wb_values[worksheet.name]

    cells: list[CellData] = []
    for row in ws_formulas.iter_rows(min_row=1, max_row=ws_formulas.max_row, max_col=ws_formulas.max_column):
        for cell in row:
            if not cell_has_signal(cell):
                continue
            is_formula = cell.data_type == "f"
            calculated_value = (
                ws_values.cell(row=cell.row, column=cell.column).value if is_formula else None
            )
            border = cell.border
            cells.append(
                CellData(
                    row=cell.row,
                    column=cell.column,
                    value=None if is_formula else cell.value,
                    formula=cell.value if is_formula else None,
                    calculated_value=calculated_value,
                    number_format=cell.number_format,
                    bold=bool(cell.font.bold),
                    italic=bool(cell.font.italic),
                    font_color=color_to_hex(cell.font.color),
                    # Only "solid" actually paints fgColor as a flat background in Excel's own
                    # rendering. Any other fill_type Excel and openpyxl left non-None — most
                    # commonly "gray125", the legacy default marker OOXML silently writes onto
                    # a cell that was merely touched by formatting (e.g. borders) without an
                    # explicit fill — would otherwise get its (often grey/black) fgColor painted
                    # as a real background here, which is not what the workbook shows on screen.
                    fill_color=color_to_hex(cell.fill.fgColor)
                    if cell.fill and cell.fill.fill_type == "solid"
                    else None,
                    horizontal_alignment=cell.alignment.horizontal,
                    vertical_alignment=cell.alignment.vertical,
                    borders={
                        "top": border.top.style if border.top else None,
                        "bottom": border.bottom.style if border.bottom else None,
                        "left": border.left.style if border.left else None,
                        "right": border.right.style if border.right else None,
                    },
                )
            )

    merged_cells = [str(cell_range) for cell_range in ws_formulas.merged_cells.ranges]
    column_widths = {
        letter: dim.width for letter, dim in ws_formulas.column_dimensions.items() if dim.width
    }
    row_heights = {
        index: dim.height for index, dim in ws_formulas.row_dimensions.items() if dim.height
    }

    result = WorksheetData(
        id=worksheet.id,
        name=worksheet.name,
        max_row=ws_formulas.max_row,
        max_column=ws_formulas.max_column,
        cells=cells,
        merged_cells=merged_cells,
        column_widths=column_widths,
        row_heights=row_heights,
    )
    print(f"[PERF] read_worksheet_data: TOTAL end-to-end: {time.perf_counter() - request_start:.3f}s", flush=True)
    return result


def apply_edits(
    db: Session, *, workbook: Workbook, worksheet: Worksheet, edits: list[dict]
) -> Worksheet:
    # TEMPORARY: total end-to-end timer for the live latency investigation — see excel_io.py's
    # _perf_log for the per-phase breakdown this should sum to. Remove both once diagnosed.
    request_start = time.perf_counter()
    path = Path(workbook.storage_path)
    # Held for the whole load-mutate-save cycle, not just the save — two overlapping edits to
    # the same workbook (or an edit racing a child-sheet create/sync) corrupted a real file
    # live by each independently reading and rewriting it at once. See
    # excel_io.workbook_write_lock()'s own docstring for the full story.
    with excel_io.workbook_write_lock(path):
        wb = excel_io.load_workbook(path, data_only=False)
        try:
            ws = wb[worksheet.name]
            cell_editor.apply_cell_edits(ws, edits)
            # The cell(s) this edit actually touched are excluded from cache restoration — see
            # save_workbook_preserving_formula_cache()'s own docstring for why reapplying their
            # *pre-edit* cached value would be wrong (most importantly when the edit changed the
            # formula itself).
            edited_coordinates = {
                (worksheet.name, ws.cell(row=edit["row"], column=edit["column"]).coordinate)
                for edit in edits
            }
            excel_io.save_workbook_preserving_formula_cache(wb, path, exclude=edited_coordinates)
        finally:
            wb.close()

    worksheet.content_updated_at = datetime.now(timezone.utc)
    db.commit()
    db.refresh(worksheet)
    print(f"[PERF] apply_edits: TOTAL end-to-end: {time.perf_counter() - request_start:.3f}s", flush=True)
    return worksheet


def delete_worksheet(db: Session, *, workbook: Workbook, worksheet: Worksheet) -> None:
    """Deletes a whole worksheet — detected frontend-side from Univer's own
    sheet.mutation.remove-sheet command (see UniverSheetGrid.tsx), the same way structural
    row/column edits are.

    If this worksheet is a *parent* in one or more child-sheet relationships, those
    relationships are gone once it's deleted (a relationship is meaningless with no parent to
    sync from) — but the *child* worksheets themselves are deliberately left alone, now
    ordinary, disconnected worksheets holding whatever data they last had. This is the
    explicit product decision here (the user's own words): "if the parent looses the data,
    the child should normally too [for column deletion]... [but for deleting the whole parent
    sheet] the child sheet is now orphan, the data in it is just static" — the frontend is
    responsible for warning about this *before* calling here, since by the time this runs the
    deletion is meant to already be confirmed.

    No app-level code has to hunt down and delete those relationship rows itself:
    sheet_relationships.parent_worksheet_id/child_worksheet_id both have a real DB-level
    ON DELETE CASCADE (see worksheet_repository.delete's own comment), so deleting this
    worksheet row removes exactly the relationship rows referencing it, on either side, and
    nothing else.
    """
    path = Path(workbook.storage_path)
    with excel_io.workbook_write_lock(path):
        wb = excel_io.load_workbook(path, data_only=False)
        try:
            if worksheet.name in wb.sheetnames:
                if len(wb.sheetnames) <= 1:
                    raise HTTPException(
                        status_code=status.HTTP_400_BAD_REQUEST,
                        detail="Cannot delete the only worksheet in a workbook",
                    )
                del wb[worksheet.name]
            excel_io.save_workbook_preserving_formula_cache(wb, path)
        finally:
            wb.close()

    worksheet_repository.delete(db, worksheet)


def apply_structural_edit(
    db: Session,
    *,
    workbook: Workbook,
    worksheet: Worksheet,
    operation: str,
    start_index: int,
    count: int,
) -> Worksheet:
    """Inserting/deleting a row or column via Univer's own UI used to be sent to the backend
    as a giant batch of per-cell value edits (reconstructing the shift as a diff) — that
    silently corrupted merged cells (a shifted edit lands on a MergedCell's read-only .value)
    and had no way to represent the operation at all beyond plain cell values. This applies
    the real structural operation directly via openpyxl instead, detected frontend-side from
    Univer's own command service (see UniverSheetGrid.tsx) rather than inferred from a diff.
    """
    path = Path(workbook.storage_path)
    # Held for the whole load-mutate-save cycle — same reasoning as apply_edits().
    with excel_io.workbook_write_lock(path):
        wb = excel_io.load_workbook(path, data_only=False)
        try:
            ws = wb[worksheet.name]
            # Unmerge everything *before* the shift, not after: a merge's non-anchor cells are
            # placeholder MergedCell objects living in ws._cells at their own coordinates, and
            # insert/delete_rows/cols moves *everything* in ws._cells (confirmed by reading its
            # source — Worksheet._move_cells has no special case for them) including those
            # placeholders. Computing "the new range" and calling unmerge_cells(old_coord)
            # afterward hits a stale coordinate whose placeholder has already been moved out
            # from under it. Unmerging first (while coordinates still match reality), then
            # shifting, then re-merging at the already-computed new coordinates sidesteps that.
            #
            # openpyxl's own unmerge_cells() still isn't safe to call here though — confirmed
            # live against this app's real data: it unconditionally does `del
            # self._cells[(row, col)]` for every non-anchor cell, but a merge read from a real
            # Excel-authored file can cover a cell that never had an XML <c> entry at all (a
            # genuinely empty cell inside the merge) — openpyxl's reader doesn't backfill a
            # MergedCell placeholder for those the way ws.merge_cells() would if called
            # programmatically, so the naive delete throws a bare KeyError.
            # filter_engine.safe_unmerge below does the same thing openpyxl's own method does,
            # just tolerating an entry that was never there to begin with (shared with
            # filter_engine.write_rows(), which needs the identical tolerance for the same
            # reason when re-writing a child sheet's own merges on sync).
            old_merged_ranges = [str(cell_range) for cell_range in ws.merged_cells.ranges]
            for coord in old_merged_ranges:
                filter_engine.safe_unmerge(ws, coord)

            if operation == "insert_row":
                ws.insert_rows(start_index, count)
            elif operation == "remove_row":
                ws.delete_rows(start_index, count)
            elif operation == "insert_col":
                ws.insert_cols(start_index, count)
            elif operation == "remove_col":
                ws.delete_cols(start_index, count)

            _remerge_shifted_ranges(
                ws=ws, old_ranges=old_merged_ranges, operation=operation, start_index=start_index, count=count
            )
            # Every cell on *this* sheet just shifted position — see
            # save_workbook_preserving_formula_cache's skip_sheets docstring for why its
            # formula caches can't be safely restored by coordinate after that. Every other
            # sheet in the workbook is untouched by this operation and still safe to restore.
            excel_io.save_workbook_preserving_formula_cache(wb, path, skip_sheets={worksheet.name})
        finally:
            wb.close()

    worksheet.content_updated_at = datetime.now(timezone.utc)
    _shift_relationships_for_structural_edit(
        db,
        parent_worksheet_id=worksheet.id,
        operation=operation,
        start_index=start_index,
        count=count,
    )
    db.commit()
    db.refresh(worksheet)
    return worksheet


def _shift_bound_after_deletion(bound: int, deleted_from: int, deleted_to: int, count: int) -> int:
    if deleted_from <= bound <= deleted_to:
        return deleted_from  # this exact row/column was itself deleted; clamp rather than guess
    if bound > deleted_to:
        return bound - count
    return bound


def _remerge_shifted_ranges(
    *, ws, old_ranges: list[str], operation: str, start_index: int, count: int
) -> None:
    """Re-applies each of `old_ranges` (already unmerged, and the sheet already
    inserted/deleted) at its shifted position. See apply_structural_edit's own comment for
    why unmerging has to happen *before* the structural edit, not after."""
    is_row_op = operation in ("insert_row", "remove_row")
    is_insert = operation in ("insert_row", "insert_col")
    deleted_to = start_index + count - 1

    def shift(bound: int) -> int:
        if is_insert:
            return bound + count if bound >= start_index else bound
        return _shift_bound_after_deletion(bound, start_index, deleted_to, count)

    for coord in old_ranges:
        min_col, min_row, max_col, max_row = range_boundaries(coord)
        if is_row_op:
            min_row, max_row = shift(min_row), shift(max_row)
        else:
            min_col, max_col = shift(min_col), shift(max_col)

        if min_row == max_row and min_col == max_col:
            continue  # collapsed to a single cell by a deletion — nothing left to merge
        ws.merge_cells(start_row=min_row, start_column=min_col, end_row=max_row, end_column=max_col)


def _remap_filter_criteria_columns(node: dict, remap: Callable[[int], int | None]) -> dict:
    """Walks a filter_criteria tree (see filter_engine.evaluate()'s docstring for its shape),
    remapping each leaf condition's column index — or dropping the condition entirely when
    `remap` returns None (the column it referenced was just deleted from the parent)."""
    if "logic" in node:
        remapped_children = [_remap_filter_criteria_columns(child, remap) for child in node.get("conditions", [])]
        return {**node, "conditions": [child for child in remapped_children if child is not None]}
    new_column = remap(node["column"])
    if new_column is None:
        return None
    return {**node, "column": new_column}


def _shift_relationships_for_structural_edit(
    db: Session, *, parent_worksheet_id: uuid.UUID, operation: str, start_index: int, count: int
) -> None:
    """Keeps every child-sheet relationship's stored positions in sync with a structural edit
    to its parent — otherwise selected_columns/header rows would silently keep pointing at
    whatever *used to* be at those positions. A column deleted from the parent is
    deliberately *dropped* from selected_columns (and any filter condition referencing it),
    not left dangling — the user's own call: child sheets should lose that data too, the same
    way the parent did. This only updates the relationship's own stored positions; the child
    worksheet's actual file is untouched here — content_updated_at's bump above already makes
    sync_service.is_outdated() report "Changes available", so the existing manual Synchronize
    flow is what actually applies the drop to the child sheet's data, same as any other
    parent edit.
    """
    relationships = sheet_relationship_repository.list_by_parent_worksheet_id(db, parent_worksheet_id)
    if not relationships:
        return

    deleted_from = start_index
    deleted_to = start_index + count - 1

    for relationship in relationships:
        if operation == "insert_row":
            if relationship.header_start_row >= start_index:
                relationship.header_start_row += count
            if relationship.header_end_row >= start_index:
                relationship.header_end_row += count
        elif operation == "remove_row":
            relationship.header_start_row = _shift_bound_after_deletion(
                relationship.header_start_row, deleted_from, deleted_to, count
            )
            relationship.header_end_row = _shift_bound_after_deletion(
                relationship.header_end_row, deleted_from, deleted_to, count
            )
            relationship.header_end_row = max(relationship.header_end_row, relationship.header_start_row)
        elif operation == "insert_col":
            relationship.selected_columns = [
                idx + count if idx >= start_index else idx for idx in relationship.selected_columns
            ]
            relationship.filter_criteria = _remap_filter_criteria_columns(
                relationship.filter_criteria, lambda idx: idx + count if idx >= start_index else idx
            )
        elif operation == "remove_col":
            relationship.selected_columns = [
                idx - count if idx > deleted_to else idx
                for idx in relationship.selected_columns
                if not (deleted_from <= idx <= deleted_to)
            ]
            relationship.filter_criteria = _remap_filter_criteria_columns(
                relationship.filter_criteria,
                lambda idx: None if deleted_from <= idx <= deleted_to else (idx - count if idx > deleted_to else idx),
            )


def list_columns(
    *, workbook: Workbook, worksheet: Worksheet, header_start_row: int, header_end_row: int
) -> list[WorksheetColumn]:
    """Lists every column in a worksheet's header block for a caller (the child-sheet creation
    form) that needs to let a human pick columns by a readable label, while the actual
    identifier used everywhere else (filtering, selection, storage) stays the 1-indexed column
    number — see filter_engine.read_rows()'s docstring for why header text alone can't safely
    be a column's identity on a real multi-row-header matrix sheet."""
    path = Path(workbook.storage_path)
    wb = excel_io.load_workbook_cached(path, data_only=True)
    ws = wb[worksheet.name]
    header_grid, _ = filter_engine.read_rows(ws, header_start_row, header_end_row)

    max_col = len(header_grid[0]) if header_grid else 0
    columns = []
    for col in range(1, max_col + 1):
        # Join each header row's resolved value for this column, skipping blanks and
        # collapsing immediate repeats — a vertically-merged single-row header (e.g. "Action"
        # spanning two header rows) would otherwise resolve to "Action" in every header row and
        # show up as the redundant-looking "Action - Action".
        label_parts: list[str] = []
        for header_row in header_grid:
            value = header_row[col - 1]
            text = "" if value is None else str(value).strip()
            if text and (not label_parts or label_parts[-1] != text):
                label_parts.append(text)
        letter = get_column_letter(col)
        columns.append(
            WorksheetColumn(index=col, letter=letter, label=" - ".join(label_parts) or f"Column {letter}")
        )
    return columns
