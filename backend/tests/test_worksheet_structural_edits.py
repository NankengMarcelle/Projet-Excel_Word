from fastapi.testclient import TestClient
from openpyxl import Workbook

from app.services.worksheet_service import _safe_unmerge


def _upload_sample(api_client: TestClient, headers: dict, sample_xlsx_bytes: bytes) -> dict:
    response = api_client.post(
        "/workbooks",
        headers=headers,
        files={"file": ("sample.xlsx", sample_xlsx_bytes, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
    )
    assert response.status_code == 201
    return response.json()


def _get_cell(cells: list[dict], row: int, column: int) -> dict:
    return next(c for c in cells if c["row"] == row and c["column"] == column)


# Same fixture layout as test_child_sheets_and_sync.py: "Data" sheet, single-row header
# (Name=1, Status=2, Amount=3), rows 2-4 data, row 5 has a merged "Total" cell (A5:B5) and a
# `=SUM(C2:C4)` formula in C5.
_CHILD_SHEET_PAYLOAD = {
    "child_sheet_name": "Active Employees",
    "header_start_row": 1,
    "header_end_row": 1,
    "selected_columns": [1, 3],
    "filter_criteria": {
        "logic": "AND",
        "conditions": [{"column": 2, "operator": "equals", "value": "Active"}],
    },
}


def _upload_and_get_ids(api_client: TestClient, headers: dict, sample_xlsx_bytes: bytes) -> tuple[str, str]:
    created = _upload_sample(api_client, headers, sample_xlsx_bytes)
    workbook_id = created["id"]
    worksheet_id = api_client.get(f"/workbooks/{workbook_id}", headers=headers).json()["worksheets"][0]["id"]
    return workbook_id, worksheet_id


def test_safe_unmerge_tolerates_a_non_anchor_cell_with_no_placeholder(tmp_path):
    """Regression test for a real crash found live against actual production data:
    openpyxl's own unmerge_cells() unconditionally does `del self._cells[(row, col)]` for
    every non-anchor cell in a merge, but a merge read from a real Excel-authored file can
    cover a cell that never had a placeholder there at all (a genuinely empty cell within the
    merge that Excel itself never wrote an XML <c> entry for) — that throws a bare KeyError.
    _safe_unmerge must tolerate this instead of crashing.
    """
    wb = Workbook()
    ws = wb.active
    ws["A1"] = "Title"
    ws.merge_cells("A1:C1")
    # Simulates the real-file condition directly: B1's MergedCell placeholder is missing,
    # exactly as if it had never been written to _cells in the first place.
    del ws._cells[(1, 2)]

    _safe_unmerge(ws, "A1:C1")

    assert "A1:C1" not in ws.merged_cells
    assert ws["A1"].value == "Title"


def test_insert_row_shifts_content_down_without_crashing_on_the_merged_cell(
    api_client: TestClient, auth_headers: dict, sample_xlsx_bytes: bytes
):
    workbook_id, worksheet_id = _upload_and_get_ids(api_client, auth_headers, sample_xlsx_bytes)

    # Inserting above row 5 (the merged "Total" row) is exactly the shape that crashed the old
    # value-diff-based autosave (a shifted cell edit landing on a MergedCell's read-only
    # .value) — this endpoint performs the real structural operation instead.
    response = api_client.patch(
        f"/workbooks/{workbook_id}/worksheets/{worksheet_id}/structure",
        headers=auth_headers,
        json={"operation": "insert_row", "start_index": 2, "count": 1},
    )
    assert response.status_code == 200

    data = api_client.get(f"/workbooks/{workbook_id}/worksheets/{worksheet_id}", headers=auth_headers).json()
    # Alice used to be row 2, now shifted to row 3.
    assert _get_cell(data["cells"], row=3, column=1)["value"] == "Alice"
    # The merged "Total" row shifted from row 5 to row 6.
    assert "A6:B6" in data["merged_cells"]
    assert _get_cell(data["cells"], row=6, column=3)["formula"] is not None


def test_remove_col_shifts_and_cascades_child_sheet_selection(
    api_client: TestClient, auth_headers: dict, sample_xlsx_bytes: bytes
):
    workbook_id, worksheet_id = _upload_and_get_ids(api_client, auth_headers, sample_xlsx_bytes)

    create_response = api_client.post(
        f"/workbooks/{workbook_id}/child-sheets",
        headers=auth_headers,
        json={"parent_worksheet_id": worksheet_id, **_CHILD_SHEET_PAYLOAD},
    )
    assert create_response.status_code == 201
    relationship_id = create_response.json()["relationship"]["id"]
    assert create_response.json()["relationship"]["selected_columns"] == [1, 3]

    # Delete "Status" (column 2) — not itself selected, but Amount (column 3) needs to shift
    # down to 2, and the filter condition referencing column 2 needs to be dropped entirely
    # (the column it filtered on no longer exists).
    response = api_client.patch(
        f"/workbooks/{workbook_id}/worksheets/{worksheet_id}/structure",
        headers=auth_headers,
        json={"operation": "remove_col", "start_index": 2, "count": 1},
    )
    assert response.status_code == 200

    relationships = api_client.get(f"/workbooks/{workbook_id}/child-sheets", headers=auth_headers).json()
    relationship = next(r for r in relationships if r["id"] == relationship_id)
    assert relationship["selected_columns"] == [1, 2]
    assert relationship["filter_criteria"]["conditions"] == []


def test_remove_col_cascades_a_selected_column_deletion_to_the_child_sheet(
    api_client: TestClient, auth_headers: dict, sample_xlsx_bytes: bytes
):
    workbook_id, worksheet_id = _upload_and_get_ids(api_client, auth_headers, sample_xlsx_bytes)

    create_response = api_client.post(
        f"/workbooks/{workbook_id}/child-sheets",
        headers=auth_headers,
        json={"parent_worksheet_id": worksheet_id, **_CHILD_SHEET_PAYLOAD},
    )
    relationship_id = create_response.json()["relationship"]["id"]

    # Delete "Name" (column 1) — it IS in selected_columns ([1, 3]). The parent lost that
    # column, so the child sheet's selection should lose it too, not keep pointing at
    # whatever now sits at position 1 (per the user's own call: "if the parent looses the
    # data, the child should normally too").
    response = api_client.patch(
        f"/workbooks/{workbook_id}/worksheets/{worksheet_id}/structure",
        headers=auth_headers,
        json={"operation": "remove_col", "start_index": 1, "count": 1},
    )
    assert response.status_code == 200

    relationships = api_client.get(f"/workbooks/{workbook_id}/child-sheets", headers=auth_headers).json()
    relationship = next(r for r in relationships if r["id"] == relationship_id)
    # Only Amount (was 3, now 2) survives; Name is gone.
    assert relationship["selected_columns"] == [2]

    status_response = api_client.get(
        f"/workbooks/{workbook_id}/child-sheets/{relationship_id}/status", headers=auth_headers
    )
    assert status_response.json()["is_outdated"] is True


def test_insert_col_shifts_child_sheet_selection_up(
    api_client: TestClient, auth_headers: dict, sample_xlsx_bytes: bytes
):
    workbook_id, worksheet_id = _upload_and_get_ids(api_client, auth_headers, sample_xlsx_bytes)

    create_response = api_client.post(
        f"/workbooks/{workbook_id}/child-sheets",
        headers=auth_headers,
        json={"parent_worksheet_id": worksheet_id, **_CHILD_SHEET_PAYLOAD},
    )
    relationship_id = create_response.json()["relationship"]["id"]

    # Insert a new column before everything — every existing column shifts right by 1.
    response = api_client.patch(
        f"/workbooks/{workbook_id}/worksheets/{worksheet_id}/structure",
        headers=auth_headers,
        json={"operation": "insert_col", "start_index": 1, "count": 1},
    )
    assert response.status_code == 200

    relationships = api_client.get(f"/workbooks/{workbook_id}/child-sheets", headers=auth_headers).json()
    relationship = next(r for r in relationships if r["id"] == relationship_id)
    assert relationship["selected_columns"] == [2, 4]
    assert relationship["filter_criteria"]["conditions"][0]["column"] == 3
