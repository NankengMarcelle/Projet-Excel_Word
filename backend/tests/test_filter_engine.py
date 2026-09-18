from openpyxl import Workbook

from app.spreadsheet.filter_engine import apply_filter, project_columns, read_rows, write_rows


def _build_multi_row_header_sheet():
    # Mirrors the real shape this feature targets: a "Prévision 2026" group header (row 1)
    # merged horizontally across two sub-columns with their own labels underneath (row 2),
    # next to a simple column whose single label is merged vertically across both header rows.
    wb = Workbook()
    ws = wb.active
    ws["A1"] = "Action"
    ws.merge_cells("A1:A2")  # vertical-only merge: label lives only in the anchor (A1)
    ws["B1"] = "Prévision 2026"
    ws.merge_cells("B1:C1")  # horizontal merge: should propagate to both B and C
    ws["B2"] = "AE voté"
    ws["C2"] = "CP voté"

    ws["A3"] = "Task one"
    ws["B3"] = 10
    ws["C3"] = 20
    ws["A4"] = "Task two"
    ws["B4"] = 30
    ws["C4"] = 40
    return ws


def test_read_rows_propagates_horizontally_merged_header_values():
    ws = _build_multi_row_header_sheet()
    header_grid, _ = read_rows(ws, header_start_row=1, header_end_row=2)

    # Row 1 (index 0): the horizontal merge (B1:C1) propagates "Prévision 2026" into both B
    # and C — a naive read would only see it in B and None in C.
    assert header_grid[0] == ["Action", "Prévision 2026", "Prévision 2026"]
    # Row 2 (index 1): A2 is part of a *vertical-only* merge anchored at A1, so it resolves to
    # "Action" too (not blank) — the merge's value is still "there" for any cell inside it.
    assert header_grid[1] == ["Action", "AE voté", "CP voté"]


def test_read_rows_keys_data_rows_by_column_index_not_header_text():
    ws = _build_multi_row_header_sheet()
    _, data_rows = read_rows(ws, header_start_row=1, header_end_row=2)

    assert data_rows == [
        {1: "Task one", 2: 10, 3: 20},
        {1: "Task two", 2: 30, 3: 40},
    ]


def test_filter_and_project_by_column_index():
    ws = _build_multi_row_header_sheet()
    _, data_rows = read_rows(ws, header_start_row=1, header_end_row=2)

    filtered = apply_filter(
        data_rows,
        {"logic": "AND", "conditions": [{"column": 2, "operator": "equals", "value": 30}]},
    )
    assert filtered == [{1: "Task two", 2: 30, 3: 40}]

    projected = project_columns(filtered, [1, 3])
    assert projected == [["Task two", 40]]


def test_write_rows_preserves_the_full_multi_row_header_block(tmp_path):
    ws = _build_multi_row_header_sheet()
    header_grid, data_rows = read_rows(ws, header_start_row=1, header_end_row=2)
    filtered = apply_filter(data_rows, {})
    projected = project_columns(filtered, [1, 3])

    wb = Workbook()
    dest = wb.active
    write_rows(dest, header_grid, [1, 3], projected)

    rows = list(dest.iter_rows(values_only=True))
    # Both header rows survive, projected down to just columns 1 and 3 (Action, CP voté) —
    # not flattened into one row of composite names.
    assert rows[0] == ("Action", "Prévision 2026")
    assert rows[1] == ("Action", "CP voté")
    assert rows[2] == ("Task one", 20)
    assert rows[3] == ("Task two", 40)


def test_read_rows_resolves_vertically_merged_data_cells_for_kept_rows():
    # A "Structure" label merged down across several task rows must still resolve to its real
    # value for a row that isn't the merge's own anchor row — otherwise a filter that keeps a
    # non-anchor row would silently show a blank where Excel visually shows a value.
    wb = Workbook()
    ws = wb.active
    ws["A1"] = "Structure"
    ws["B1"] = "Task"
    ws["A2"] = "CIRT"
    ws["B2"] = "Task one"
    ws["B3"] = "Task two"
    ws.merge_cells("A2:A3")  # "CIRT" visually applies to both row 2 and row 3

    _, data_rows = read_rows(ws, header_start_row=1, header_end_row=1)
    assert data_rows == [
        {1: "CIRT", 2: "Task one"},
        {1: "CIRT", 2: "Task two"},
    ]
