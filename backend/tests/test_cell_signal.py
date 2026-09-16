from openpyxl import Workbook
from openpyxl.styles import Alignment, Font

from app.spreadsheet.cell_signal import cell_has_signal


def test_plain_untouched_cell_has_no_signal():
    wb = Workbook()
    ws = wb.active
    cell = ws.cell(row=1, column=1)
    assert cell_has_signal(cell) is False


def test_black_font_color_is_not_signal():
    # Regression test for a real bug found in a real production workbook: a blanket
    # formatting sweep across ~500 columns wrote font_color="FF000000" (black — the default
    # text color) onto every cell it touched, with no actual data. Treating that as "real
    # content" made used_range() report the sheet's full 525-declared-column width instead of
    # its real ~18-column extent, defeating the whole point of trimming to used content.
    wb = Workbook()
    ws = wb.active
    cell = ws.cell(row=1, column=1)
    cell.font = Font(color="FF000000")
    assert cell_has_signal(cell) is False


def test_non_black_font_color_is_signal():
    wb = Workbook()
    ws = wb.active
    cell = ws.cell(row=1, column=1)
    cell.font = Font(color="FFC81E1E")
    assert cell_has_signal(cell) is True


def test_bottom_vertical_alignment_is_not_signal():
    # Same category of bug as the font-color case above, on Excel's default vertical
    # alignment for an unwrapped cell.
    wb = Workbook()
    ws = wb.active
    cell = ws.cell(row=1, column=1)
    cell.alignment = Alignment(vertical="bottom")
    assert cell_has_signal(cell) is False


def test_top_vertical_alignment_is_signal():
    wb = Workbook()
    ws = wb.active
    cell = ws.cell(row=1, column=1)
    cell.alignment = Alignment(vertical="top")
    assert cell_has_signal(cell) is True


def test_general_horizontal_alignment_is_not_signal():
    wb = Workbook()
    ws = wb.active
    cell = ws.cell(row=1, column=1)
    cell.alignment = Alignment(horizontal="general")
    assert cell_has_signal(cell) is False


def test_center_horizontal_alignment_is_signal():
    wb = Workbook()
    ws = wb.active
    cell = ws.cell(row=1, column=1)
    cell.alignment = Alignment(horizontal="center")
    assert cell_has_signal(cell) is True


def test_bold_cell_is_signal_even_with_default_color_and_alignment():
    # The presence of one real signal (bold) shouldn't be masked by also having
    # default-equivalent values on the other attributes.
    wb = Workbook()
    ws = wb.active
    cell = ws.cell(row=1, column=1)
    cell.font = Font(bold=True, color="FF000000")
    cell.alignment = Alignment(vertical="bottom")
    assert cell_has_signal(cell) is True


def test_a_value_is_always_signal_regardless_of_style():
    wb = Workbook()
    ws = wb.active
    cell = ws.cell(row=1, column=1, value="x")
    assert cell_has_signal(cell) is True
