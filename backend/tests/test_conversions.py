import io

from docx import Document
from fastapi.testclient import TestClient


def _upload_sample(api_client: TestClient, headers: dict, sample_xlsx_bytes: bytes) -> dict:
    response = api_client.post(
        "/workbooks",
        headers=headers,
        files={"file": ("sample.xlsx", sample_xlsx_bytes, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
    )
    assert response.status_code == 201
    return response.json()


def test_convert_worksheet_and_download_repeatedly(
    api_client: TestClient, auth_headers: dict, sample_xlsx_bytes: bytes
):
    created = _upload_sample(api_client, auth_headers, sample_xlsx_bytes)
    workbook_id = created["id"]
    worksheet_id = api_client.get(
        f"/workbooks/{workbook_id}", headers=auth_headers
    ).json()["worksheets"][0]["id"]

    convert_response = api_client.post(f"/worksheets/{worksheet_id}/convert", headers=auth_headers)
    assert convert_response.status_code == 201
    body = convert_response.json()
    conversion_id = body["conversion"]["id"]
    assert body["word_document"]["filename"] == "Data.docx"

    meta_response = api_client.get(f"/conversions/{conversion_id}", headers=auth_headers)
    assert meta_response.status_code == 200

    download_response = api_client.get(f"/conversions/{conversion_id}/download", headers=auth_headers)
    assert download_response.status_code == 200

    document = Document(io.BytesIO(download_response.content))
    table = document.tables[0]
    header_texts = [cell.text for cell in table.rows[0].cells]
    assert header_texts == ["Name", "Status", "Amount"]
    assert table.rows[0].cells[0].paragraphs[0].runs[0].bold is True

    # Converted Word documents stay available in the Word Files page indefinitely, downloadable
    # on demand — a second (or third) download must succeed exactly like the first, not 410.
    second_download = api_client.get(f"/conversions/{conversion_id}/download", headers=auth_headers)
    assert second_download.status_code == 200
    assert second_download.content == download_response.content


def test_list_conversions_reflects_download_state_and_ownership(
    api_client: TestClient, auth_headers: dict, sample_xlsx_bytes: bytes
):
    from tests.conftest import register_and_login

    created = _upload_sample(api_client, auth_headers, sample_xlsx_bytes)
    workbook_id = created["id"]
    worksheet = api_client.get(f"/workbooks/{workbook_id}", headers=auth_headers).json()["worksheets"][0]

    convert_response = api_client.post(f"/worksheets/{worksheet['id']}/convert", headers=auth_headers)
    conversion_id = convert_response.json()["conversion"]["id"]

    listed = api_client.get("/conversions", headers=auth_headers).json()
    assert len(listed) == 1
    row = listed[0]
    assert row["conversion_id"] == conversion_id
    assert row["filename"] == "Data.docx"
    assert row["worksheet_name"] == worksheet["name"]
    assert row["workbook_filename"] == "sample.xlsx"
    assert row["downloaded_at"] is None

    api_client.get(f"/conversions/{conversion_id}/download", headers=auth_headers)
    listed_after_download = api_client.get("/conversions", headers=auth_headers).json()
    assert listed_after_download[0]["downloaded_at"] is not None

    _, other_token = register_and_login(api_client)
    other_headers = {"Authorization": f"Bearer {other_token}"}
    assert api_client.get("/conversions", headers=other_headers).json() == []


def test_delete_conversion_removes_it_and_requires_ownership(
    api_client: TestClient, auth_headers: dict, sample_xlsx_bytes: bytes
):
    from tests.conftest import register_and_login

    created = _upload_sample(api_client, auth_headers, sample_xlsx_bytes)
    workbook_id = created["id"]
    worksheet_id = api_client.get(
        f"/workbooks/{workbook_id}", headers=auth_headers
    ).json()["worksheets"][0]["id"]

    conversion_id = api_client.post(
        f"/worksheets/{worksheet_id}/convert", headers=auth_headers
    ).json()["conversion"]["id"]

    _, other_token = register_and_login(api_client)
    other_headers = {"Authorization": f"Bearer {other_token}"}
    assert api_client.delete(f"/conversions/{conversion_id}", headers=other_headers).status_code == 404

    delete_response = api_client.delete(f"/conversions/{conversion_id}", headers=auth_headers)
    assert delete_response.status_code == 204

    assert api_client.get("/conversions", headers=auth_headers).json() == []
    assert api_client.get(f"/conversions/{conversion_id}", headers=auth_headers).status_code == 404
    assert api_client.get(f"/conversions/{conversion_id}/download", headers=auth_headers).status_code == 404


def test_convert_worksheet_requires_ownership(
    api_client: TestClient, auth_headers: dict, sample_xlsx_bytes: bytes
):
    from tests.conftest import register_and_login

    created = _upload_sample(api_client, auth_headers, sample_xlsx_bytes)
    workbook_id = created["id"]
    worksheet_id = api_client.get(
        f"/workbooks/{workbook_id}", headers=auth_headers
    ).json()["worksheets"][0]["id"]

    _, other_token = register_and_login(api_client)
    other_headers = {"Authorization": f"Bearer {other_token}"}

    response = api_client.post(f"/worksheets/{worksheet_id}/convert", headers=other_headers)
    assert response.status_code == 404
