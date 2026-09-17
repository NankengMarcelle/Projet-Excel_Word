from datetime import date

import pytest
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from app.spreadsheet.cell_signal import used_range
from app.spreadsheet.word_exporter import _format_number, worksheet_to_docx


def test_used_range_ignores_phantom_dimension_beyond_real_content():
    # Regression test for the real bug: a sheet with a 3x3 block of actual data, but whose
    # *declared* dimensions balloon far beyond that because a distant cell was formatted
    # (no value, just a border) — the same pattern found in a real production workbook
    # (one sheet: 528x525 declared, only 7,395 of 277,200 cells with any real content, all
    # within the first 18 columns). used_range() must report the real 3x3 extent, not the
    # phantom one, or the Word export ends up sized to the phantom range instead.
    wb = Workbook()
    ws = wb.active
    for row in range(1, 4):
        for col in range(1, 4):
            ws.cell(row=row, column=col, value=f"r{row}c{col}")
    # Just touching the cell (no value, no style) is enough for openpyxl to register it and
    # inflate ws.max_row/ws.max_column — giving it a border here would be wrong: a border IS
    # real signal by cell_has_signal's own rules, so used_range would correctly include it.
    # The real bug is cells that inflate the declared dimensions *without* tripping any
    # signal check at all, which is what this reproduces.
    ws.cell(row=500, column=300)

    assert ws.max_row == 500
    assert ws.max_column == 300

    assert used_range(ws) == (3, 3)


def test_worksheet_to_docx_trims_to_real_content(tmp_path):
    wb = Workbook()
    ws = wb.active
    ws.cell(row=1, column=1, value="Name")
    ws.cell(row=2, column=1, value="Alice")
    ws.cell(row=500, column=300)  # touched, unstyled — see the test above for why

    output_path = tmp_path / "out.docx"
    worksheet_to_docx(ws, output_path)

    document = Document(output_path)
    table = document.tables[0]
    # Would be 500 rows x 300 columns (150,000 mostly-blank cells) without the used_range()
    # fix — this is the actual "2+ minute conversion, 57-page document" bug reproduced at
    # unit-test scale.
    assert len(table.rows) == 2
    assert len(table.columns) == 1


def test_worksheet_to_docx_does_not_shade_gray125_placeholder_fill(tmp_path):
    # "gray125" is the legacy OOXML marker left on a cell that was merely touched by
    # formatting (e.g. a border) without an explicit fill — it must never be painted as a
    # real background, or every such cell comes out looking (usually grey/black) filled in
    # the Word doc when Excel itself shows it as plain white.
    wb = Workbook()
    ws = wb.active
    cell = ws.cell(row=1, column=1, value="X")
    cell.fill = PatternFill(fill_type="gray125")

    output_path = tmp_path / "out.docx"
    worksheet_to_docx(ws, output_path)

    document = Document(output_path)
    doc_cell = document.tables[0].cell(0, 0)
    shd = doc_cell._tc.find(".//{http://schemas.openxmlformats.org/wordprocessingml/2006/main}shd")
    assert shd is None


def test_worksheet_to_docx_shades_solid_fill(tmp_path):
    wb = Workbook()
    ws = wb.active
    cell = ws.cell(row=1, column=1, value="X")
    cell.fill = PatternFill(start_color="FFFF00", end_color="FFFF00", fill_type="solid")

    output_path = tmp_path / "out.docx"
    worksheet_to_docx(ws, output_path)

    document = Document(output_path)
    doc_cell = document.tables[0].cell(0, 0)
    shd = doc_cell._tc.find(".//{http://schemas.openxmlformats.org/wordprocessingml/2006/main}shd")
    assert shd is not None
    assert shd.get("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}fill") == "FFFF00"


def test_worksheet_to_docx_formats_dates_readably(tmp_path):
    wb = Workbook()
    ws = wb.active
    cell = ws.cell(row=1, column=1, value=date(2026, 9, 16))
    cell.number_format = "YYYY-MM-DD"

    output_path = tmp_path / "out.docx"
    worksheet_to_docx(ws, output_path)

    document = Document(output_path)
    assert document.tables[0].cell(0, 0).text == "2026-09-16"


def test_worksheet_to_docx_applies_vertical_alignment(tmp_path):
    from openpyxl.styles import Alignment

    wb = Workbook()
    ws = wb.active
    cell = ws.cell(row=1, column=1, value="X")
    cell.alignment = Alignment(vertical="center")

    output_path = tmp_path / "out.docx"
    worksheet_to_docx(ws, output_path)

    document = Document(output_path)
    from docx.enum.table import WD_ALIGN_VERTICAL

    assert document.tables[0].cell(0, 0).vertical_alignment == WD_ALIGN_VERTICAL.CENTER


def test_worksheet_to_docx_applies_font_color_without_crashing(tmp_path):
    # Regression test for a real crash hit live (not synthetic): color_to_hex() returns
    # openpyxl's raw 8-char ARGB string, but RGBColor.from_string() raises outright on
    # anything but 6 hex characters — every colored-font cell 500'd the whole conversion
    # (`ValueError: RGBColor() takes three integer values 0-255`), which the frontend showed
    # as a Convert button stuck on "Converting..." forever rather than a visible error.
    wb = Workbook()
    ws = wb.active
    cell = ws.cell(row=1, column=1, value="X")
    cell.font = Font(color="FFC81E1E")  # ARGB, as openpyxl stores it

    output_path = tmp_path / "out.docx"
    worksheet_to_docx(ws, output_path)

    document = Document(output_path)
    run = document.tables[0].cell(0, 0).paragraphs[0].runs[0]
    assert run.font.color.rgb == RGBColor.from_string("C81E1E")


# --- Number format rendering ------------------------------------------------------------


def test_format_number_percentage():
    assert _format_number(0.15, "0%") == "15%"
    assert _format_number(0.1534, "0.00%") == "15.34%"


def test_format_number_thousands_and_decimals():
    assert _format_number(1234.5, "#,##0.00") == "1,234.50"
    assert _format_number(1234, "#,##0") == "1,234"


def test_format_number_currency():
    assert _format_number(1234.5, "$#,##0.00") == "$1,234.50"


def test_format_number_returns_none_for_general_or_non_numeric():
    assert _format_number(5, "General") is None
    assert _format_number("text", "#,##0.00") is None
    assert _format_number(True, "#,##0.00") is None  # bool is technically an int subclass


def test_format_number_ignores_conditional_color_and_multi_section_formats():
    # "[Red]-#,##0;[Blue]#,##0" style multi-section formats — only the first (positive)
    # section should be interpreted, with the [Red]/[Blue] conditional-color tags stripped.
    assert _format_number(1234, "[Blue]#,##0;[Red]-#,##0") == "1,234"


def test_format_number_ignores_spacer_currency_symbols():
    # Regression test, found from a real report: `_-* #,##0\ _€_-` is a French accounting
    # format for a *plain whole number*, no currency at all — the "€" only appears inside a
    # `_€` spacer escape (Excel's "reserve blank space the width of this character, but don't
    # actually show it" directive), used purely so this column visually aligns with real
    # currency columns next to it. Excel itself never displays that €; treating it as "this
    # cell is currency" prefixed every value with a € sign that shouldn't have been there.
    assert _format_number(1234, r"_-* #,##0\ _€_-") == "1,234"


def test_worksheet_to_docx_renders_percentage_cell(tmp_path):
    wb = Workbook()
    ws = wb.active
    cell = ws.cell(row=1, column=1, value=0.42)
    cell.number_format = "0%"

    output_path = tmp_path / "out.docx"
    worksheet_to_docx(ws, output_path)

    document = Document(output_path)
    assert document.tables[0].cell(0, 0).text == "42%"


# --- Border style/color fidelity ---------------------------------------------------------


def test_worksheet_to_docx_preserves_border_style_and_color(tmp_path):
    # Regression test: every border used to render identically as a thin black single line
    # regardless of the source's actual style/color/thickness.
    wb = Workbook()
    ws = wb.active
    cell = ws.cell(row=1, column=1, value="X")
    cell.border = Border(top=Side(style="thick", color="FFFF0000"))

    output_path = tmp_path / "out.docx"
    worksheet_to_docx(ws, output_path)

    document = Document(output_path)
    doc_cell = document.tables[0].cell(0, 0)
    top = doc_cell._tc.find(
        ".//{http://schemas.openxmlformats.org/wordprocessingml/2006/main}tcBorders"
        "/{http://schemas.openxmlformats.org/wordprocessingml/2006/main}top"
    )
    ns = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
    assert top.get(f"{ns}val") == "single"
    assert top.get(f"{ns}sz") == "24"  # "thick" -> 24 eighths-of-a-point
    assert top.get(f"{ns}color") == "FF0000"


def test_worksheet_to_docx_defaults_border_color_when_unset(tmp_path):
    wb = Workbook()
    ws = wb.active
    cell = ws.cell(row=1, column=1, value="X")
    cell.border = Border(top=Side(style="thin"))  # no color specified

    output_path = tmp_path / "out.docx"
    worksheet_to_docx(ws, output_path)

    document = Document(output_path)
    doc_cell = document.tables[0].cell(0, 0)
    ns = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
    top = doc_cell._tc.find(f".//{ns}tcBorders/{ns}top")
    assert top.get(f"{ns}color") == "000000"


# --- Fit-to-page: columns and font shrink together, like Excel's own print scaling ------


def test_worksheet_to_docx_shrinks_font_proportionally_to_fit_wide_sheet(tmp_path):
    # Regression test: the first fix tried was shrinking column widths alone (to fit many
    # columns on one landscape page) while leaving font size untouched — that made long labels
    # wrap letter-by-letter instead of word-by-word, confirmed live against a real workbook.
    # The actual fix mirrors Excel's own "Fit to page width" print scaling: shrink column
    # widths *and* font size together by the same factor, so nothing has to wrap mid-word.
    wb = Workbook()
    ws = wb.active
    # 5 columns at width 20 sums to a good deal more than one landscape page can hold, but not
    # so much that scaling would clamp both fonts down to the floor (which would flatten out
    # the very proportionality this test exists to check) — see the floor-clamp test below for
    # that extreme case instead.
    for col in range(1, 6):
        ws.column_dimensions[get_column_letter(col)].width = 20
    cell_with_explicit_size = ws.cell(row=1, column=1, value="Explicit")
    cell_with_explicit_size.font = Font(size=14)
    cell_with_default_size = ws.cell(row=1, column=2, value="Default")  # no Font() override at all
    # used_range() only extends to columns with real cell *content* — a column-width-only
    # column with no value in it wouldn't count, so give the last column a value too, or this
    # test's sheet would only "really" be 2 columns wide regardless of the widths set above.
    ws.cell(row=1, column=5, value="Last")

    output_path = tmp_path / "out.docx"
    worksheet_to_docx(ws, output_path)

    document = Document(output_path)
    table = document.tables[0]
    explicit_run = table.cell(0, 0).paragraphs[0].runs[0]
    default_run = table.cell(0, 1).paragraphs[0].runs[0]

    # Both shrank...
    assert explicit_run.font.size.pt < 14
    assert default_run.font.size.pt < 11
    # ...by the *same* factor, so their relative sizes are preserved (14pt was always meant to
    # look larger than the 11pt default, before or after scaling).
    assert explicit_run.font.size.pt / default_run.font.size.pt == pytest.approx(14 / 11, rel=0.01)
    # 5 columns at 20 Excel-width-units each is wider than any standard page — column widths
    # must have actually shrunk too, not just gotten set to their naive (large) size.
    assert table.columns[0].width < Inches(2)


def test_worksheet_to_docx_floors_font_size_rather_than_shrinking_to_nothing(tmp_path):
    # An extreme column count/width combination shouldn't drive the font down to an
    # unreadable sliver (or zero) — it should stop at a legible floor instead.
    wb = Workbook()
    ws = wb.active
    for col in range(1, 16):
        ws.column_dimensions[get_column_letter(col)].width = 30
    cell = ws.cell(row=1, column=1, value="X")
    cell.font = Font(size=14)
    ws.cell(row=1, column=15, value="Last")

    output_path = tmp_path / "out.docx"
    worksheet_to_docx(ws, output_path)

    document = Document(output_path)
    run = document.tables[0].cell(0, 0).paragraphs[0].runs[0]
    assert run.font.size.pt == pytest.approx(6.0)


def test_worksheet_to_docx_does_not_shrink_font_when_sheet_already_fits(tmp_path):
    wb = Workbook()
    ws = wb.active
    cell = ws.cell(row=1, column=1, value="X")
    cell.font = Font(size=12)

    output_path = tmp_path / "out.docx"
    worksheet_to_docx(ws, output_path)

    document = Document(output_path)
    run = document.tables[0].cell(0, 0).paragraphs[0].runs[0]
    assert run.font.size == Pt(12)


# --- Merged cells shouldn't accumulate blank paragraphs ---------------------------------


def test_merging_a_wide_range_does_not_leave_blank_trailing_paragraphs(tmp_path):
    # Regression test for a real bug found from a live screenshot comparison, not guessed:
    # python-docx's Cell.merge() concatenates every merged cell's own paragraphs onto the
    # anchor's. Every cell in this table starts with one (empty) paragraph even if never
    # written to, so merging N cells left N-1 blank trailing paragraphs stacked after the
    # anchor's real text — a single tight title line in the source came out as real text
    # followed by 7+ blank lines in Word, a large, obviously-wrong visual gap.
    wb = Workbook()
    ws = wb.active
    ws.cell(row=1, column=1, value="Title")
    # Columns B-F: real cells, but never written to (no value) — exactly what a wide
    # full-row title merge looks like in a real report.
    ws.merge_cells("A1:F1")

    output_path = tmp_path / "out.docx"
    worksheet_to_docx(ws, output_path)

    document = Document(output_path)
    doc_cell = document.tables[0].cell(0, 0)
    assert len(doc_cell.paragraphs) == 1
    assert doc_cell.text == "Title"


def test_merging_preserves_a_real_embedded_newline_in_the_anchor_cell(tmp_path):
    # The fix isn't "keep only the first paragraph" — trimming has to be relative to however
    # many paragraphs the anchor cell itself legitimately had *before* merging, not a fixed
    # "1". (In practice python-docx's own .text setter renders an embedded "\n" as a line
    # break within a single paragraph rather than a second paragraph — confirmed directly —
    # so today that's always 1 for this app's cells; this test pins that behavior so a future
    # python-docx version — or a future change to how cell text gets set — that starts
    # splitting on "\n" into real paragraphs doesn't silently regress back to the original
    # bug via a trim that assumes exactly 1.)
    wb = Workbook()
    ws = wb.active
    ws.cell(row=1, column=1, value="Line one\nLine two")
    ws.merge_cells("A1:D1")

    output_path = tmp_path / "out.docx"
    worksheet_to_docx(ws, output_path)

    document = Document(output_path)
    doc_cell = document.tables[0].cell(0, 0)
    assert doc_cell.text == "Line one\nLine two"


# --- Formula cells fall back to formula text when the cached value is missing -----------


def test_worksheet_to_docx_falls_back_to_formula_text_when_cached_value_is_missing(tmp_path):
    # openpyxl never computes formulas itself, so a workbook it just created — or one whose
    # cache was lost through an edit-save round trip (see
    # excel_io.save_workbook_preserving_formula_cache's own docstring for how that happens on
    # a real workbook) — has no cached value for a formula cell at all. Without a fallback,
    # data_only=True gives back None with no way to tell "this was a formula" from "this cell
    # is genuinely empty", and the cell renders completely blank.
    wb = Workbook()
    ws = wb.active
    sheet_name = ws.title
    ws["A1"] = 5
    ws["A2"] = "=A1*2"
    saved_path = tmp_path / "source.xlsx"
    wb.save(saved_path)
    wb.close()

    ws_values = load_workbook(saved_path, data_only=True)[sheet_name]
    ws_formulas = load_workbook(saved_path, data_only=False)[sheet_name]
    assert ws_values["A2"].value is None  # confirms the premise: no cached value at all

    output_path = tmp_path / "out.docx"
    worksheet_to_docx(ws_values, output_path, ws_formulas=ws_formulas)

    document = Document(output_path)
    table = document.tables[0]
    assert table.cell(0, 0).text == "5"
    assert table.cell(1, 0).text == "=A1*2"


def test_worksheet_to_docx_without_a_formulas_view_still_renders_blank(tmp_path):
    # Existing behavior when the caller has no second (data_only=False) load handy — every
    # other test in this file relies on this staying exactly as it was. B2 (a plain value)
    # is what pulls row 2 into the exported range at all here — a bare, cacheless formula
    # cell with no formatting of its own wouldn't register as real content in the values-only
    # view (see used_range()'s call site in worksheet_to_docx for why), so without it this
    # test's row wouldn't exist to check in the first place.
    wb = Workbook()
    ws = wb.active
    sheet_name = ws.title
    ws["A2"] = "=A1*2"
    ws["B2"] = "unrelated"
    saved_path = tmp_path / "source.xlsx"
    wb.save(saved_path)
    wb.close()

    ws_values = load_workbook(saved_path, data_only=True)[sheet_name]

    output_path = tmp_path / "out.docx"
    worksheet_to_docx(ws_values, output_path)

    document = Document(output_path)
    table = document.tables[0]
    assert table.cell(1, 0).text == ""
    assert table.cell(1, 1).text == "unrelated"
