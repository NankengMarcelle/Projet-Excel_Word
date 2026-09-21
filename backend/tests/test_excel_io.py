import threading
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font

from app.core.config import settings
from app.spreadsheet import excel_io, object_storage


def _seed_cached_formula_value(path, sheet_name: str, coordinate: str, value) -> None:
    # openpyxl itself never computes formulas, so a workbook it just created has no cached
    # value at all for a formula cell — this stands in for "a real spreadsheet application
    # already computed and saved this file at some point", which is the actual starting state
    # every real workbook this app receives is in.
    excel_io._reinject_formula_cache(path, {(sheet_name, coordinate): value})


def _build_workbook_with_cross_sheet_formula(tmp_path):
    path = tmp_path / "wb.xlsx"
    wb = Workbook()
    sheet1 = wb.active
    sheet1.title = "Sheet1"
    sheet1["A1"] = 5
    sheet2 = wb.create_sheet("Sheet2")
    sheet2["A1"] = "=Sheet1!A1*2"
    wb.save(path)
    wb.close()
    _seed_cached_formula_value(path, "Sheet2", "A1", 10)
    return path


def test_preserves_untouched_formula_cached_value_across_an_unrelated_edit(tmp_path):
    # Regression test for the actual bug found live: apply_edits() loads with
    # data_only=False (needed to keep formula *text* intact for editing), which never holds
    # a formula's cached *value* in the first place — an ordinary save through that path
    # silently blanked every formula cell's displayed value workbook-wide, not just whatever
    # was actually edited.
    path = _build_workbook_with_cross_sheet_formula(tmp_path)

    wb = load_workbook(path, data_only=False)
    wb["Sheet1"]["A1"].font = Font(bold=True)  # an edit that never touches the formula cell
    excel_io.save_workbook_preserving_formula_cache(wb, path, exclude=set())
    wb.close()

    values = load_workbook(path, data_only=True)
    assert values["Sheet2"]["A1"].value == 10
    assert values["Sheet1"]["A1"].value == 5

    formulas = load_workbook(path, data_only=False)
    assert formulas["Sheet2"]["A1"].value == "=Sheet1!A1*2"
    assert formulas["Sheet1"]["A1"].font.bold is True


def test_does_not_reapply_a_stale_value_to_the_cell_that_was_actually_edited(tmp_path):
    # A cell's *own* pre-edit cached value belongs to whatever it held before this edit — most
    # importantly, reapplying it would be actively wrong if the edit changed the formula
    # itself, since the old cached number no longer corresponds to the new formula at all.
    path = _build_workbook_with_cross_sheet_formula(tmp_path)

    wb = load_workbook(path, data_only=False)
    wb["Sheet2"]["A1"] = "=Sheet1!A1*3"  # a different formula than the one cached at 10
    excel_io.save_workbook_preserving_formula_cache(wb, path, exclude={("Sheet2", "A1")})
    wb.close()

    formulas = load_workbook(path, data_only=False)
    assert formulas["Sheet2"]["A1"].value == "=Sheet1!A1*3"

    values = load_workbook(path, data_only=True)
    # No formula engine recomputed this, and the stale "10" (for the *old* formula) was
    # correctly withheld — None, not a wrong number, is the honest result here.
    assert values["Sheet2"]["A1"].value is None


def test_workbook_write_lock_serializes_concurrent_saves_without_corrupting_the_file(tmp_path):
    # Regression test for a real bug: a production workbook was corrupted live when two
    # requests (an autosave PUT and a child-sheet creation) each ran their own
    # load-mutate-save cycle for the same file concurrently, with nothing serializing "load,
    # mutate, save" as one unit. Their writes interleaved and garbled the .xlsx's zip
    # directory structure — confirmed by hand: all of the file's individual entries were
    # still intact, only the central directory was corrupted, exactly what two overlapping
    # zipfile writers to the same path produce. workbook_write_lock() closes that window.
    path = tmp_path / "wb.xlsx"
    wb = Workbook()
    wb.active["A1"] = 0
    wb.save(path)
    wb.close()

    errors: list[Exception] = []

    def bump():
        try:
            with excel_io.workbook_write_lock(path):
                w = excel_io.load_workbook(path, data_only=False)
                try:
                    w.active["A1"] = (w.active["A1"].value or 0) + 1
                    excel_io.save_workbook(w, path)
                finally:
                    w.close()
        except Exception as exc:  # pragma: no cover - failure path, asserted on below
            errors.append(exc)

    threads = [threading.Thread(target=bump) for _ in range(20)]
    for thread in threads:
        thread.start()
    for thread in threads:
        thread.join()

    assert errors == []
    # The file must still be a valid, loadable workbook after 20 concurrent load-mutate-save
    # cycles — and, since workbook_write_lock() serializes them, every single increment must
    # have landed (a race would either corrupt the file outright or silently lose updates).
    result = load_workbook(path)
    assert result.active["A1"].value == 20


# --- Remote storage mirror (STORAGE_BACKEND="s3") ---------------------------
#
# These monkeypatch object_storage's download/upload/delete with in-memory fakes rather than
# hitting a real S3-compatible endpoint — the point is confirming excel_io.py calls them at the
# right moments, not exercising boto3 itself.


def test_load_workbook_downloads_from_remote_on_local_cache_miss(tmp_path, monkeypatch):
    monkeypatch.setattr(settings, "STORAGE_BACKEND", "s3")
    monkeypatch.setattr(settings, "STORAGE_ROOT", str(tmp_path))
    downloads: list[str] = []

    def fake_download(key: str, local_path: Path) -> None:
        downloads.append(key)
        local_path.parent.mkdir(parents=True, exist_ok=True)
        Workbook().save(local_path)

    monkeypatch.setattr(object_storage, "download", fake_download)

    missing_path = tmp_path / "workbooks" / "owner" / "wb.xlsx"
    wb = excel_io.load_workbook(missing_path)
    wb.close()

    assert downloads == ["workbooks/owner/wb.xlsx"]


def test_load_workbook_does_not_redownload_once_present_locally(tmp_path, monkeypatch):
    # Single-process/single-instance assumption (see workbook_write_lock's docstring): once a
    # file has been fetched once in this process's lifetime, every writer re-uploads
    # immediately, so the local copy stays authoritative — no per-request network round trip.
    monkeypatch.setattr(settings, "STORAGE_BACKEND", "s3")
    monkeypatch.setattr(settings, "STORAGE_ROOT", str(tmp_path))
    path = tmp_path / "wb.xlsx"
    wb = Workbook()
    wb.save(path)
    wb.close()

    downloads: list[str] = []
    monkeypatch.setattr(object_storage, "download", lambda key, local_path: downloads.append(key))

    wb = excel_io.load_workbook(path)
    wb.close()

    assert downloads == []


def test_save_workbook_uploads_to_remote(tmp_path, monkeypatch):
    monkeypatch.setattr(settings, "STORAGE_BACKEND", "s3")
    monkeypatch.setattr(settings, "STORAGE_ROOT", str(tmp_path))
    uploads: list[str] = []
    monkeypatch.setattr(object_storage, "upload", lambda local_path, key: uploads.append(key))

    path = tmp_path / "workbooks" / "owner" / "wb.xlsx"
    wb = Workbook()
    excel_io.save_workbook(wb, path)
    wb.close()

    assert uploads == ["workbooks/owner/wb.xlsx"]


def test_delete_object_deletes_from_remote(tmp_path, monkeypatch):
    monkeypatch.setattr(settings, "STORAGE_BACKEND", "s3")
    monkeypatch.setattr(settings, "STORAGE_ROOT", str(tmp_path))
    path = tmp_path / "workbooks" / "owner" / "wb.xlsx"
    path.parent.mkdir(parents=True)
    path.write_bytes(b"x")

    deletes: list[str] = []
    monkeypatch.setattr(object_storage, "delete", lambda key: deletes.append(key))

    excel_io.delete_object(path)

    assert not path.exists()
    assert deletes == ["workbooks/owner/wb.xlsx"]


def test_local_backend_never_touches_object_storage(tmp_path, monkeypatch):
    # Regression guard: the default ("local") backend must be a complete no-op for every
    # object_storage function, so existing local-dev behavior is untouched by this module.
    calls: list[str] = []
    monkeypatch.setattr(object_storage, "download", lambda *a, **k: calls.append("download"))
    monkeypatch.setattr(object_storage, "upload", lambda *a, **k: calls.append("upload"))
    monkeypatch.setattr(object_storage, "delete", lambda *a, **k: calls.append("delete"))

    path = tmp_path / "wb.xlsx"
    wb = Workbook()
    excel_io.save_workbook(wb, path)
    wb.close()
    excel_io.load_workbook(path).close()
    excel_io.delete_object(path)

    assert calls == []
