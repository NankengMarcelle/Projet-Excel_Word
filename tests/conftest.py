import io
import uuid

import pytest
from fastapi.testclient import TestClient
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

from app.main import app

client = TestClient(app)


@pytest.fixture
def api_client() -> TestClient:
    return client


def _build_sample_workbook_bytes() -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Data"

    headers = ["Name", "Status", "Amount"]
    for col, header in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.font = Font(bold=True)
        cell.fill = PatternFill(start_color="FFFF00", end_color="FFFF00", fill_type="solid")

    rows = [("Alice", "Active", 100), ("Bob", "Inactive", 200), ("Carol", "Active", 300)]
    for row_index, (name, status, amount) in enumerate(rows, start=2):
        ws.cell(row=row_index, column=1, value=name)
        ws.cell(row=row_index, column=2, value=status)
        amount_cell = ws.cell(row=row_index, column=3, value=amount)
        amount_cell.number_format = "#,##0.00"

    ws.cell(row=5, column=2, value="Total")
    ws.cell(row=5, column=3, value="=SUM(C2:C4)")
    ws.merge_cells("A5:B5")
    ws.column_dimensions["A"].width = 20

    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


@pytest.fixture
def sample_xlsx_bytes() -> bytes:
    return _build_sample_workbook_bytes()


def register_and_login(api_client: TestClient) -> tuple[str, str]:
    email = f"user-{uuid.uuid4().hex[:10]}@example.com"
    password = "correct-horse-battery-staple"
    api_client.post(
        "/auth/register", json={"email": email, "password": password, "full_name": "Test User"}
    )
    login_response = api_client.post("/auth/login", data={"username": email, "password": password})
    token = login_response.json()["access_token"]
    return email, token


@pytest.fixture
def auth_headers(api_client: TestClient) -> dict:
    _, token = register_and_login(api_client)
    return {"Authorization": f"Bearer {token}"}
